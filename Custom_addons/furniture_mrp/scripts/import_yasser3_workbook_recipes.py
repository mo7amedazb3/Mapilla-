# -*- coding: utf-8 -*-
"""One-shot, audited recipe replacement for database ``yasser3``.

Run through ``odoo-bin shell``.  The script deliberately builds the new
recipes inactive, validates the whole batch, archives the old active BoMs,
then activates the new batch in the same database transaction.

Only the newer, complete workbook is used for recipe values.  The second
workbook is an older four-sheet subset and is retained in the cold backup for
audit purposes.
"""

import json
import os
import re
from collections import Counter, defaultdict

from openpyxl import load_workbook
from odoo.exceptions import ValidationError


SOURCE_XLSX = os.environ.get(
    'FURNITURE_RECIPE_XLSX',
    '/root/.codex/attachments/2f31789c-f6d9-461f-bab2-d265a8fec775/'
    'تكاليف استاذ ياسر.xlsx',
)
COMMIT = os.environ.get('FURNITURE_IMPORT_COMMIT') == '1'
VALIDATION_DATABASE = os.environ.get('FURNITURE_IMPORT_VALIDATION_DB', '')


STAGE_ORDER = (
    'priming',
    'painting',
    'carpentry',
    'bases',
    'finishing',
    'tailoring',
    'upholstery',
    'packaging',
)

STAGE_KEYWORDS = (
    ('تقديم', 'priming'),
    ('دهانات', 'painting'),
    ('تجميع', 'carpentry'),
    ('قواعد', 'bases'),
    ('تجهيز', 'finishing'),
    ('قماش', 'tailoring'),
    ('تفصيل', 'tailoring'),
    ('كسوه', 'upholstery'),
    ('كسوة', 'upholstery'),
    ('تغليف', 'packaging'),
)

UOM_ALIASES = {
    'متر': 'meter',
    'متر مكعب': 'cubic_meter',
    'كيلو': 'kg',
    'كيلو جرام': 'kg',
    'لوح': 'board',
    'فرخ': 'sheet',
    'لفة': 'roll',
    'علبه': 'box',
    'علبة': 'box',
    'قطعه': 'unit',
    'قطعة': 'unit',
    'توب': 'bolt',
}

MATERIAL_ALIASES = {
    'co10m': 'CO 10m',
    'co 10m': 'CO 10m',
    'co16m': 'CO 16m',
    'co 16m': 'CO 16m',
    'mdf3m': 'MDF 3m',
    'mdf 3m': 'MDF 3m',
    'mdf7m': 'MDF 7m',
    'mdf 7m': 'MDF 7m',
    'mdf10m': 'MDF 10m',
    'mdf10 m': 'MDF 10m',
    'mdf 10m': 'MDF 10m',
    'خشب زان 5': 'خشب زان 5 سم',
    'خشب بياض 4.4': 'خشب بياض 4.4 سم',
    'مسمار 4 هوا': 'مسمار 4 سم هواء',
    'سنفره': 'فرخ سنفره 36 مقاس صغير',
    'اجو': 'آجو رش',
    'عبك شبح': 'عبك شبح خفيف',
    'واير': 'واير زجزاج',
    'اسفنج .5 سم': 'اسفنج 5 مم كـ18',
    'اسفنج .8 سم': 'اسفنج 8 مم كـ22',
    'اسفنج 1.8 سم': 'اسفنج1.8 سم كـ22',
    'اسفنج 5 سم كـ33': 'اسفنج 5 سم هارد كـ33',
    'اسفنج 7 سم سوفت': 'اسفنج 7 سم سوبر سوفت',
    'اسفنج 10 سم كـ30سوفت': 'اسفنج 10 سم سوبر سوفت كـ30',
    'اسفنج 10 سم كـ35سوفت': 'اسفنج 10 سم سوفت كـ35',
    'اسفنج 4.5 سم سوفت': 'اسفنج 4.5 سم سوبر سوفت',
    'استيك اخضر': 'أستيك اخضر عرض 5سم',
    'استيك اصفر': 'أستيك اصفر عرض 6سم',
    'فايبر هواو': 'فايبر هولو',
    'مسامير 13سم': 'مسامير 13 سم',
}

PRODUCT_ALIASES = {
    'كنبه كبيره': 'كنبة كبيرة',
    'كنبه صغيره': 'كنبة صغيرة',
    'صغيره': 'كنبة صغيرة',
    'شيز لونج': 'شازلونج',
    'شيزلونج': 'شازلونج',
    'شيزلونج': 'شازلونج',
    'زاويه': 'زاوية',
    'وحده': 'وحدة',
}


def clean(value):
    if value is None:
        return ''
    return ' '.join(str(value).replace('\u00a0', ' ').split()).strip()


def key(value):
    value = clean(value).casefold()
    value = value.replace('ـ', '')
    value = value.replace('أ', 'ا').replace('إ', 'ا').replace('آ', 'ا')
    value = value.replace('ة', 'ه').replace('ى', 'ي')
    value = re.sub(r'\s+', ' ', value)
    return value.strip()


MATERIAL_ALIASES = {key(alias): target for alias, target in MATERIAL_ALIASES.items()}
PRODUCT_ALIASES = {key(alias): target for alias, target in PRODUCT_ALIASES.items()}
UOM_ALIASES = {key(alias): target for alias, target in UOM_ALIASES.items()}

# The newer workbook accidentally leaves the UoM blank for Wavy's yellow
# elastic, while the same SKU is explicitly measured by the bolt in both
# Oafly and Mony.  Infer that one value from the two corroborating sheets;
# genuinely unitless sundries remain an audited unit fallback.
MISSING_UOM_INFERENCE = {
    key('استيك اصفر'): 'توب',
}

# The live database contains two historic, stock-used SKUs with this exact
# normalized name.  ID 65 is the older and more frequently referenced one;
# pinning it makes the migration deterministic without merging stock history.
PREFERRED_EXISTING_PRODUCT_IDS = {
    key('مسامير 13 سم'): 65,
}

# These are the two generic placeholders used by the per-order tailoring
# allocation screen.  Classifying them ensures a manager's chosen fabric and
# takawe replace the placeholder quantities instead of being added twice.
TAILORING_MATERIAL_KINDS = {
    key('قماش'): 'fabric',
    key('مشجر'): 'takawe',
}


def number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.replace(',', '').strip())
        except ValueError:
            return None
    return None


def stage_from_heading(value, row=None):
    text = key(value)
    for needle, code in STAGE_KEYWORDS:
        if key(needle) in text:
            return code
    # The complete source workbook has a literal "F" instead of the second
    # stage heading in Beig Moon.  Its row position is unambiguous.
    if row == 14 and text == 'f':
        return 'carpentry'
    return None


def canonical_products(labels):
    result = []
    for label in labels:
        normalized = PRODUCT_ALIASES.get(key(label), clean(label))
        if normalized and normalized not in result:
            result.append(normalized)
    return result


def canonical_material(name, stage, uom_code):
    raw_key = key(name)
    if raw_key == key('كرتون') and (
        stage == 'packaging' or uom_code == 'kg'
    ):
        return 'كرتون تغليف'
    return MATERIAL_ALIASES.get(raw_key, clean(name))


def workbook_blocks(ws):
    if ws.max_column > 15:
        return [
            {'heading_col': 2, 'name_cols': (3,), 'material_col': 3,
             'uom_col': 4, 'qty_col': 5, 'price_col': 6, 'kit_qty': 1.0},
            {'heading_col': 9, 'name_cols': (10,), 'material_col': 10,
             'uom_col': 11, 'qty_col': 12, 'price_col': 13, 'kit_qty': 1.0},
            {'heading_col': 16, 'name_cols': (17,), 'material_col': 17,
             'uom_col': 18, 'qty_col': 19, 'price_col': 20, 'kit_qty': 1.0},
            {'heading_col': 23, 'name_cols': (24,), 'material_col': 24,
             'uom_col': 25, 'qty_col': 26, 'price_col': 27, 'kit_qty': 1.0},
        ]
    return [
        {'heading_col': 2, 'name_cols': (4, 5), 'material_col': 4,
         'uom_col': 5, 'qty_col': 6, 'price_col': 7, 'kit_qty': 1.0},
        {'heading_col': 2, 'name_cols': (10,), 'material_col': 10,
         'uom_col': 11, 'qty_col': 12, 'price_col': 13, 'kit_qty': 2.0},
    ]


def parse_workbook(path):
    wb = load_workbook(path, data_only=False, read_only=False)
    models = []
    issues = []
    for ws in wb.worksheets:
        model_name = clean(ws.title)
        recipes = []
        for block_index, cfg in enumerate(workbook_blocks(ws), 1):
            labels = [
                clean(ws.cell(4, col).value)
                for col in cfg['name_cols']
                if clean(ws.cell(4, col).value)
            ]
            label_keys = {key(label) for label in labels}
            product_names = canonical_products(labels)
            if not product_names:
                issues.append({
                    'kind': 'missing_output', 'sheet': ws.title,
                    'block': block_index,
                })
                continue
            stage = None
            lines = []
            for row in range(5, ws.max_row + 1):
                heading = ws.cell(row, cfg['heading_col']).value
                detected_stage = stage_from_heading(heading, row=row)
                if detected_stage:
                    stage = detected_stage
                    continue
                item = clean(ws.cell(row, cfg['material_col']).value)
                if not item or not stage:
                    continue
                item_key = key(item)
                item_without_count = re.sub(
                    r'^\d+(?:[.,]\d+)?\s*', '', item_key,
                ).strip()
                if (
                    item_key in {'البند', 'الفئه'}
                    or item_key in label_keys
                    or item_without_count in label_keys
                    or 'اجمالي' in item_key
                    or item_key.startswith('=')
                ):
                    continue
                qty_value = number(ws.cell(row, cfg['qty_col']).value)
                uom_raw = clean(ws.cell(row, cfg['uom_col']).value)
                if not uom_raw and item_key in MISSING_UOM_INFERENCE:
                    uom_raw = MISSING_UOM_INFERENCE[item_key]
                    issues.append({
                        'kind': 'inferred_uom', 'sheet': ws.title,
                        'block': block_index, 'row': row, 'item': item,
                        'uom': uom_raw,
                    })
                uom_code = UOM_ALIASES.get(key(uom_raw), 'unit')
                price_value = number(ws.cell(row, cfg['price_col']).value)
                if qty_value is None:
                    issues.append({
                        'kind': 'missing_qty', 'sheet': ws.title,
                        'block': block_index, 'row': row, 'item': item,
                    })
                    continue
                if qty_value <= 0:
                    continue
                if not uom_raw:
                    issues.append({
                        'kind': 'defaulted_uom', 'sheet': ws.title,
                        'block': block_index, 'row': row, 'item': item,
                        'uom': 'unit',
                    })
                material_name = canonical_material(item, stage, uom_code)
                material_kind = TAILORING_MATERIAL_KINDS.get(key(material_name))
                # The workbook lists its generic patterned-fabric placeholder
                # under tailoring.  In the operational allocation screen this
                # placeholder is the takawe choice, whose authoritative stage
                # is upholstery.  Store it there from the outset so a manager's
                # selection replaces the placeholder instead of being added as
                # a second material requirement.
                effective_stage = (
                    'upholstery' if material_kind == 'takawe' else stage
                )
                lines.append({
                    'stage': effective_stage,
                    'material_raw': item,
                    'material': material_name,
                    'uom_raw': uom_raw,
                    'uom_code': uom_code,
                    'qty': qty_value,
                    'unit_price': price_value,
                    'row': row,
                })
            # Every finished-product label is independent, but each recipe must
            # retain the exact quantities printed in its workbook block.  Do
            # not divide a shared left-hand block between its side-by-side
            # products, and do not divide the armchair block by the Kit count.
            # The workbook quantities are the authoritative per-recipe values.
            for output_index, product_name in enumerate(product_names, 1):
                recipes.append({
                    'block': block_index,
                    'output': output_index,
                    'product': product_name,
                    'kit_qty': cfg['kit_qty'],
                    'lines': [
                        {
                            **line,
                            'source_qty': line['qty'],
                            'qty': line['qty'],
                        }
                        for line in lines
                    ],
                })
        models.append({
            'model': model_name,
            'kit': 'طقم %s' % model_name,
            'recipes': recipes,
        })
    return models, issues


def pick_mode(values):
    values = [float(value) for value in values if value is not None and value > 0]
    if not values:
        return 0.0
    counts = Counter(values)
    highest = max(counts.values())
    candidates = {value for value, count in counts.items() if count == highest}
    for value in reversed(values):
        if value in candidates:
            return value
    return values[-1]


def find_by_normalized_name(
    records,
    name,
    extra_filter=None,
    preferred_id=None,
):
    target = key(name)
    matches = records.browse()
    for record in records:
        if key(record.name) != target:
            continue
        if extra_filter and not extra_filter(record):
            continue
        matches |= record
    if preferred_id:
        preferred = matches.filtered(lambda record: record.id == preferred_id)
        if preferred:
            return preferred[:1]
    if len(matches) > 1:
        raise ValidationError(
            'Ambiguous existing product match for %s: %s'
            % (name, ', '.join(str(record_id) for record_id in matches.ids))
        )
    if matches:
        return matches[:1]
    return records.browse()


def build_summary(models, issues):
    recipes = [recipe for model in models for recipe in model['recipes']]
    lines = [line for recipe in recipes for line in recipe['lines']]
    return {
        'models': len(models),
        'recipes': len(recipes),
        'stage_lines': len(lines),
        'kit_lines': len(recipes),
        'missing_qty': sum(issue['kind'] == 'missing_qty' for issue in issues),
        'defaulted_uom': sum(issue['kind'] == 'defaulted_uom' for issue in issues),
        'inferred_uom': sum(issue['kind'] == 'inferred_uom' for issue in issues),
        'products': sorted({recipe['product'] for recipe in recipes}),
        'materials': len({line['material'] for line in lines}),
    }


def source_batch_material_totals(models):
    totals = defaultdict(float)
    for model_data in models:
        seen_blocks = set()
        for recipe in model_data['recipes']:
            if recipe['block'] in seen_blocks:
                continue
            seen_blocks.add(recipe['block'])
            for line in recipe['lines']:
                bucket = (
                    key(model_data['model']), line['stage'],
                    key(line['material']), line['uom_code'],
                )
                totals[bucket] += line['source_qty']
    return {bucket: round(qty, 6) for bucket, qty in totals.items()}


def kit_material_totals(models):
    totals = defaultdict(float)
    for model_data in models:
        for recipe in model_data['recipes']:
            for line in recipe['lines']:
                bucket = (
                    key(model_data['model']), line['stage'],
                    key(line['material']), line['uom_code'],
                )
                totals[bucket] += line['qty'] * recipe['kit_qty']
    return {bucket: round(qty, 6) for bucket, qty in totals.items()}


def run(env, commit=False):
    is_validation_db = bool(
        VALIDATION_DATABASE
        and VALIDATION_DATABASE.startswith('yasser3_recipe_validate_')
        and env.cr.dbname == VALIDATION_DATABASE
    )
    if env.cr.dbname != 'yasser3' and not is_validation_db:
        raise ValidationError(
            'This importer may run on yasser3 or its explicit validation clone only.'
        )
    if not os.path.isfile(SOURCE_XLSX):
        raise ValidationError('Workbook not found: %s' % SOURCE_XLSX)

    models_data, issues = parse_workbook(SOURCE_XLSX)
    summary = build_summary(models_data, issues)
    expected = {
        'models': 11,
        'recipes': 31,
        'stage_lines': 803,
        'kit_lines': 31,
        'missing_qty': 10,
        'defaulted_uom': 25,
        'inferred_uom': 1,
        'materials': 60,
    }
    for field_name, expected_value in expected.items():
        if summary[field_name] != expected_value:
            raise ValidationError(
                'Workbook preflight mismatch for %s: got %s, expected %s'
                % (field_name, summary[field_name], expected_value)
            )

    expected_model_recipes = {
        'بيج مون': [('كنبة كبيرة', 1.0), ('شازلونج', 1.0), ('فوتيه', 2.0)],
        'مارلي': [('كنبة كبيرة', 1.0), ('كنبة صغيرة', 1.0), ('فوتيه', 2.0)],
        'مورا': [('كنبة كبيرة', 1.0), ('كنبة صغيرة', 1.0), ('فوتيه', 2.0)],
        'بيبندوم': [('كنبة كبيرة', 1.0), ('فوتيه', 2.0)],
        'ويفي': [('كنبة كبيرة', 1.0), ('كنبة صغيرة', 1.0), ('فوتيه', 2.0)],
        'مارفيل': [('كنبة كبيرة', 1.0), ('فوتيه', 2.0)],
        'اوفلي': [('كنبة كبيرة', 1.0), ('فوتيه', 2.0)],
        'موفينج': [
            ('مخدع شمال / يمين', 1.0),
            ('شفيز', 1.0),
            ('زاوية', 1.0),
            ('فوتيه', 1.0),
        ],
        'موني': [('كنبة كبيرة', 1.0), ('فوتيه', 2.0)],
        'فيرال': [
            ('مخدع شمال / يمين', 1.0),
            ('شفيز', 1.0),
            ('شازلونج', 1.0),
            ('وحدة', 1.0),
        ],
        'بالهوم': [('كنبة كبيرة', 1.0), ('شازلونج', 1.0), ('فوتيه', 2.0)],
    }
    actual_model_recipes = {
        model_data['model']: [
            (recipe['product'], recipe['kit_qty'])
            for recipe in model_data['recipes']
        ]
        for model_data in models_data
    }
    if actual_model_recipes != expected_model_recipes:
        raise ValidationError(
            'Workbook product split mismatch: got %s, expected %s'
            % (actual_model_recipes, expected_model_recipes)
        )
    source_totals = source_batch_material_totals(models_data)
    if any(
        line['qty'] != line['source_qty']
        for model_data in models_data
        for recipe in model_data['recipes']
        for line in recipe['lines']
    ):
        raise ValidationError(
            'A product recipe quantity differs from its workbook source value.'
        )
    expected_kit_totals = kit_material_totals(models_data)

    print('PREFLIGHT ' + json.dumps(summary, ensure_ascii=False, sort_keys=True))
    print('ISSUES ' + json.dumps(issues, ensure_ascii=False, sort_keys=True))
    if not commit:
        env.cr.rollback()
        print('DRY_RUN_OK')
        return summary

    Model = env['furniture.product.model'].sudo().with_context(active_test=False)
    Product = env['product.product'].sudo().with_context(active_test=False)
    Bom = env['mrp.bom'].sudo().with_context(active_test=False)
    unit_uom = env.ref('uom.product_uom_unit')
    company = env.company

    all_models = Model.search([]).sorted('id')
    target_models = {}
    for sequence, model_data in enumerate(models_data, 10):
        model = find_by_normalized_name(all_models, model_data['model'])
        if model:
            model.write({'active': True, 'sequence': sequence})
        else:
            model = Model.create({
                'name': model_data['model'],
                'sequence': sequence,
                'active': True,
            })
            all_models |= model
        target_models[key(model_data['model'])] = model

    old_active_boms = Bom.search([('active', '=', True)]).filtered(
        lambda bom: bool(
            bom.furniture_product_id
            or bom.furniture_recipe_model_id
            or bom.furniture_model_id
            or bom.furniture_is_model_recipe
        )
    ).sorted('id')
    old_active_bom_ids = old_active_boms.ids
    old_active_model_ids = Model.search([('active', '=', True)]).ids

    product_prices = defaultdict(list)
    for model_data in models_data:
        for recipe in model_data['recipes']:
            for line in recipe['lines']:
                product_prices[key(line['material'])].append(line['unit_price'])

    all_products = Product.search([]).sorted('id')
    raw_products = {}
    created_raw_product_ids = []
    for material_name in sorted(
        {line['material'] for model in models_data
         for recipe in model['recipes'] for line in recipe['lines']},
        key=key,
    ):
        material_kind = TAILORING_MATERIAL_KINDS.get(key(material_name))
        product = find_by_normalized_name(
            all_products,
            material_name,
            extra_filter=lambda rec: (
                not rec.furniture_dimension_source_product_id
                and not rec.furniture_model_id
                and not rec.product_tmpl_id.furniture_model_id
                and not Bom.search_count([
                    ('furniture_product_id', '=', rec.id),
                    ('type', 'in', ('normal', 'phantom')),
                ])
            ),
            preferred_id=PREFERRED_EXISTING_PRODUCT_IDS.get(key(material_name)),
        )
        if not product:
            product = Product.create({
                'name': material_name,
                'type': 'consu',
                'is_storable': True,
                'uom_id': unit_uom.id,
                'uom_po_id': unit_uom.id,
                'sale_ok': False,
                'purchase_ok': True,
                'standard_price': pick_mode(product_prices[key(material_name)]),
                'furniture_tailoring_material_kind': material_kind or False,
                'active': True,
            })
            all_products |= product
            created_raw_product_ids.append(product.id)
        elif (
            not product.active
            or not product.is_storable
            or (
                material_kind
                and product.furniture_tailoring_material_kind != material_kind
            )
        ):
            product.write({
                'active': True,
                'is_storable': True,
                **(
                    {'furniture_tailoring_material_kind': material_kind}
                    if material_kind else {}
                ),
            })
        raw_products[key(material_name)] = product

    recipe_products = {}
    created_finished_product_ids = []
    for product_name in summary['products']:
        product = find_by_normalized_name(
            all_products,
            product_name,
            extra_filter=lambda rec: (
                not rec.furniture_dimension_source_product_id
                and not rec.furniture_model_id
                and not rec.product_tmpl_id.furniture_model_id
            ),
        )
        if not product:
            product = Product.create({
                'name': product_name,
                'type': 'consu',
                'is_storable': True,
                'uom_id': unit_uom.id,
                'uom_po_id': unit_uom.id,
                'sale_ok': True,
                'purchase_ok': False,
                'active': True,
            })
            all_products |= product
            created_finished_product_ids.append(product.id)
        elif not product.active or not product.is_storable:
            product.write({'active': True, 'is_storable': True})
        recipe_products[key(product_name)] = product

    recipes_by_product = defaultdict(list)
    for model_data in models_data:
        model = target_models[key(model_data['model'])]
        for recipe in model_data['recipes']:
            recipes_by_product[key(recipe['product'])].append((model, recipe))

    new_masters = Bom.browse()
    new_hidden = Bom.browse()
    sequence_by_stage = {stage: (index + 1) * 100 for index, stage in enumerate(STAGE_ORDER)}
    for product_key, model_recipes in recipes_by_product.items():
        product = recipe_products[product_key]
        first_model = model_recipes[0][0]
        master = Bom.with_context(
            furniture_skip_model_recipe_sync=True,
            tracking_disable=True,
            mail_create_nolog=True,
        ).create({
            'code': 'XLSX/%s' % clean(product.name),
            'product_tmpl_id': product.product_tmpl_id.id,
            'furniture_product_id': product.id,
            'product_qty': 1.0,
            'product_uom_id': product.uom_id.id,
            'type': 'normal',
            'company_id': company.id,
            'active': False,
            'furniture_recipe_model_id': first_model.id,
            'furniture_recipe_ready': True,
        })
        new_masters |= master
        for model, recipe in model_recipes:
            child_id = master.action_select_furniture_recipe_model(model.id)
            child = Bom.browse(child_id).exists()
            commands = [(5, 0, 0)]
            stage_line_index = Counter()
            for line in recipe['lines']:
                stage_line_index[line['stage']] += 1
                material = raw_products[key(line['material'])]
                material_kind = TAILORING_MATERIAL_KINDS.get(
                    key(line['material'])
                )
                commands.append((0, 0, {
                    'sequence': sequence_by_stage[line['stage']] + stage_line_index[line['stage']],
                    'stage': line['stage'],
                    'product_id': material.id,
                    'product_qty': line['qty'],
                    'product_uom_code': line['uom_code'],
                    'quantity_mode': 'fixed',
                    'furniture_tailoring_material_kind': material_kind or False,
                }))
            child.with_context(
                furniture_skip_model_recipe_sync=True,
                tracking_disable=True,
            ).write({
                'active': False,
                'furniture_recipe_ready': True,
                'use_priming': True,
                'use_painting': True,
                'use_carpentry': True,
                'use_bases': True,
                'use_finishing': True,
                'use_tailoring': True,
                'use_sewing': False,
                'use_upholstery': True,
                'use_packaging': True,
                'furniture_stage_material_line_ids': commands,
            })
            new_hidden |= child

    if len(new_masters) != len(summary['products']):
        raise ValidationError('Unexpected master recipe count.')
    if len(new_hidden) != summary['recipes']:
        raise ValidationError('Unexpected hidden recipe count.')
    if len(new_hidden.furniture_stage_material_line_ids) != summary['stage_lines']:
        raise ValidationError('Unexpected imported material line count.')
    if any(not line.product_uom_id for line in new_hidden.furniture_stage_material_line_ids):
        raise ValidationError('Imported recipe line without a UoM.')

    # Atomic cut-over: old records remain available to historic production
    # orders, but disappear from every active recipe selector.
    old_active_boms.write({'active': False})
    new_masters.write({'active': True})
    new_hidden.write({'active': True, 'furniture_recipe_ready': True})

    new_kits = Bom.browse()
    for model_data in models_data:
        model = target_models[key(model_data['model'])]
        kit_product = find_by_normalized_name(
            all_products,
            model_data['kit'],
            extra_filter=lambda rec, model=model: (
                rec.furniture_model_id == model
                or rec.product_tmpl_id.furniture_model_id == model
            ),
        )
        if not kit_product:
            kit_product = Product.create({
                'name': model_data['kit'],
                'type': 'consu',
                'is_storable': True,
                'uom_id': unit_uom.id,
                'uom_po_id': unit_uom.id,
                'sale_ok': True,
                'purchase_ok': False,
                'active': True,
            })
            all_products |= kit_product
        elif not kit_product.active or not kit_product.is_storable:
            kit_product.write({'active': True, 'is_storable': True})
        line_commands = []
        for recipe in model_data['recipes']:
            component = recipe_products[key(recipe['product'])]
            line_commands.append((0, 0, {
                'product_id': component.id,
                'product_qty': recipe['kit_qty'],
                'product_uom_id': component.uom_id.id,
            }))
        kit = Bom.create({
            'code': 'XLSX/KIT/%s' % model_data['model'],
            'product_tmpl_id': kit_product.product_tmpl_id.id,
            'furniture_product_id': kit_product.id,
            'product_qty': 1.0,
            'product_uom_id': kit_product.uom_id.id,
            'type': 'phantom',
            'company_id': company.id,
            'active': True,
            'furniture_model_id': model.id,
            'bom_line_ids': line_commands,
        })
        new_kits |= kit

    actual_db_totals = defaultdict(float)
    for kit in new_kits:
        model = kit.furniture_model_id
        for component_line in kit.bom_line_ids:
            source_component = (
                component_line.product_id.furniture_dimension_source_product_id
                or component_line.product_id
            )
            recipe = new_hidden.filtered(
                lambda child, model=model, source_component=source_component: (
                    child.furniture_model_id == model
                    and child.furniture_product_id == source_component
                )
            )
            if len(recipe) != 1:
                raise ValidationError(
                    'Kit component does not resolve to exactly one imported recipe.'
                )
            for stage_line in recipe.furniture_stage_material_line_ids:
                bucket = (
                    key(model.name), stage_line.stage,
                    key(stage_line.product_id.name),
                    stage_line.product_uom_code,
                )
                actual_db_totals[bucket] += (
                    stage_line.product_qty * component_line.product_qty
                )
    actual_db_totals = {
        bucket: round(qty, 6)
        for bucket, qty in actual_db_totals.items()
    }
    if actual_db_totals != expected_kit_totals:
        mismatches = [
            (
                bucket,
                expected_kit_totals.get(bucket),
                actual_db_totals.get(bucket),
            )
            for bucket in sorted(
                set(expected_kit_totals) | set(actual_db_totals)
            )
            if expected_kit_totals.get(bucket) != actual_db_totals.get(bucket)
        ]
        raise ValidationError(
            'Imported Kit totals do not match the full workbook recipe values: %s'
            % (mismatches[:12],)
        )

    obsolete_composite_keys = {
        key('كنبة كبيرة + شازلونج'),
        key('كنبة كبيرة + كنبة صغيرة'),
    }
    obsolete_composite_products = Product.search([
        ('active', '=', True),
    ]).filtered(
        lambda product: key(product.name) in obsolete_composite_keys
    )
    obsolete_composite_products.write({'active': False})
    if Product.search([('active', '=', True)]).filtered(
        lambda product: key(product.name) in obsolete_composite_keys
    ):
        raise ValidationError('An obsolete composite finished product is still active.')

    for old_model in Model.search([('active', '=', True)]):
        if old_model.id not in {record.id for record in target_models.values()}:
            old_model.write({'active': False})

    active_models = Model.search([('active', '=', True)])
    active_boms = Bom.search([('active', '=', True)]).filtered(
        lambda bom: bool(
            bom.furniture_product_id
            or bom.furniture_recipe_model_id
            or bom.furniture_model_id
            or bom.furniture_is_model_recipe
        )
    )
    active_hidden = active_boms.filtered('furniture_is_model_recipe')
    active_masters = active_boms.filtered(
        lambda bom: bom.type == 'normal' and not bom.furniture_is_model_recipe
    )
    active_kits = active_boms.filtered(lambda bom: bom.type == 'phantom')
    active_stage_lines = active_hidden.furniture_stage_material_line_ids
    if len(active_models) != 11:
        raise ValidationError('Active model cut-over count is not 11.')
    if set(active_masters.ids) != set(new_masters.ids):
        raise ValidationError('An old visible master recipe is still active.')
    if set(active_hidden.ids) != set(new_hidden.ids):
        raise ValidationError('An old hidden model recipe is still active.')
    if set(active_kits.ids) != set(new_kits.ids):
        raise ValidationError('An old Kit recipe is still active.')
    if len(active_stage_lines) != 803:
        raise ValidationError('Active recipe material count is not 803.')
    if len(new_kits.bom_line_ids) != 31:
        raise ValidationError('Kit component line count is not 31.')

    result = {
        **summary,
        'old_active_boms_archived': len(old_active_bom_ids),
        'old_active_models_before': len(old_active_model_ids),
        'active_models_after': len(active_models),
        'new_master_boms': len(new_masters),
        'new_hidden_boms': len(new_hidden),
        'new_kit_boms': len(new_kits),
        'created_raw_products': len(created_raw_product_ids),
        'created_finished_products': len(created_finished_product_ids),
        'source_batch_material_buckets': len(source_totals),
        'kit_material_buckets': len(expected_kit_totals),
        'full_sheet_quantities_verified': True,
        'archived_composite_products': len(obsolete_composite_products),
        'master_ids': new_masters.ids,
        'hidden_ids': new_hidden.ids,
        'kit_ids': new_kits.ids,
    }
    env.cr.commit()
    print('IMPORT_OK ' + json.dumps(result, ensure_ascii=False, sort_keys=True))
    return result


if 'env' in globals():
    run(env, commit=COMMIT)
