import base64
import importlib.util
from pathlib import Path
from unittest.mock import patch

from lxml import etree
from openpyxl import Workbook

from odoo.addons.furniture_mrp.models import mrp_tailoring_material_setup
from odoo.exceptions import AccessError, UserError, ValidationError
from odoo.tests.common import TransactionCase, new_test_user


TEST_IMAGE = (
    b'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwC'
    b'AAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII='
)
TEST_IMAGE_2 = (
    b'iVBORw0KGgoAAAANSUhEUgAAAAIAAAACCAYAAABytg0k'
    b'AAAAFElEQVR4nGNkYPj/n4GBgYGJAQoAHRkCAjRcHicAAAAASUVORK5CYII='
)


class TestTailoringMaterialSetup(TransactionCase):
    """Regression coverage for the compact per-product material setup.

    The selected fabric and takawe are order-line decisions.  They must not
    mutate the product or its BoM, and must remain stable when the ordinary
    production material cache is completely rebuilt.
    """

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.manager_user = new_test_user(
            cls.env,
            login='furniture_tailoring_setup_manager',
            groups=(
                'base.group_user,'
                'furniture_mrp.group_furniture_mrp_manager'
            ),
        )
        cls.regular_user = new_test_user(
            cls.env,
            login='furniture_tailoring_setup_regular',
            groups='base.group_user',
        )
        cls.supervisor_user = new_test_user(
            cls.env,
            login='furniture_tailoring_setup_supervisor',
            groups=(
                'base.group_user,'
                'furniture_mrp.group_furniture_mrp_supervisor_tailoring'
            ),
        )

    def _new_raw_product(self, name, kind=False, uom=False, image=False):
        uom = uom or self.env.ref('uom.product_uom_unit')
        product = self.env['product.product'].create({
            'name': '%s %s' % (name, self._testMethodName),
            'type': 'consu',
            'is_storable': True,
            'uom_id': uom.id,
            'uom_po_id': uom.id,
            'image_1920': image or False,
        })
        if kind:
            product.product_tmpl_id.write({
                'furniture_tailoring_material_kind': kind,
            })
        return product

    def _create_finished_product(self, name, model, materials, image=False):
        product = self.env['product.product'].create({
            'name': '%s %s' % (name, self._testMethodName),
            'type': 'consu',
            'is_storable': True,
            'furniture_model_id': model.id,
            'image_1920': image or False,
        })
        bom = self.env['mrp.bom'].create({
            'product_tmpl_id': product.product_tmpl_id.id,
            'product_qty': 1.0,
            'type': 'normal',
            'furniture_product_id': product.id,
            'furniture_recipe_model_id': model.id,
            'use_tailoring': True,
            'use_upholstery': True,
            'furniture_stage_material_line_ids': [
                (0, 0, {
                    'stage': stage,
                    'product_id': raw.id,
                    'product_qty': qty,
                    'product_uom_id': raw.uom_id.id,
                    'quantity_mode': 'scaled',
                })
                for raw, qty, stage in materials
            ],
        })
        product.invalidate_recordset([
            'furniture_is_finished_product',
            'furniture_has_active_normal_recipe',
        ])
        return product, bom

    def _create_fixture(self, line_quantities=(2.0, 1.0, 3.0)):
        fixture_serial = getattr(self, '_tailoring_fixture_serial', 0) + 1
        self._tailoring_fixture_serial = fixture_serial
        model = self.env['furniture.product.model'].create({
            'name': 'Tailoring Setup Model %s %s' % (
                self._testMethodName,
                fixture_serial,
            ),
        })
        buyer = self.env['res.partner'].create({
            'name': 'Tailoring Setup Buyer %s' % self._testMethodName,
            'is_company': True,
        })
        beneficiary = self.env['res.partner'].create({
            'name': 'Tailoring Setup Beneficiary %s' % self._testMethodName,
        })
        auto_fabric = self._new_raw_product('Automatic Fabric', 'fabric')
        manual_fabric = self._new_raw_product('Selected Fabric', 'fabric')
        second_fabric = self._new_raw_product('Second Selected Fabric', 'fabric')
        auto_takawe = self._new_raw_product('Automatic Takawe', 'takawe')
        manual_takawe = self._new_raw_product('Selected Takawe', 'takawe')
        other_material = self._new_raw_product('Tailoring Thread')

        products = self.env['product.product']
        boms = self.env['mrp.bom']
        for index, label in enumerate((
            'Large Sofa',
            'Small Sofa',
            'Chaise Longue',
        )):
            product, bom = self._create_finished_product(
                label,
                model,
                (
                    (auto_fabric, 2.0 + index, 'tailoring'),
                    (auto_takawe, 1.0, 'upholstery'),
                    (other_material, 0.5, 'tailoring'),
                ),
                image=TEST_IMAGE if index == 0 else False,
            )
            products |= product
            boms |= bom

        production = self.env['furniture.mrp.production'].with_context(
            furniture_skip_material_refresh=True,
            furniture_skip_stage_plan_sync=True,
        ).create({
            'product_qty': sum(line_quantities),
            'state': 'confirmed',
            'furniture_order_model_id': model.id,
            'buyer_partner_id': buyer.id,
            'beneficiary_partner_id': beneficiary.id,
            'use_tailoring': True,
            'use_upholstery': True,
        })
        lines = self.env['furniture.mrp.production.line']
        for sequence, product, bom, qty in zip(
            (10, 20, 30), products, boms, line_quantities,
        ):
            lines |= self.env['furniture.mrp.production.line'].with_context(
                furniture_skip_material_refresh=True,
                furniture_skip_stage_plan_sync=True,
            ).create({
                'production_id': production.id,
                'sequence': sequence,
                'product_id': product.id,
                'furniture_order_model_id': model.id,
                'buyer_partner_id': buyer.id,
                'beneficiary_partner_id': beneficiary.id,
                'bom_id': bom.id,
                'product_qty': qty,
                'stage_selection_initialized': True,
                'use_tailoring': True,
                'use_upholstery': True,
            })
        production._refresh_material_lines_for_stage_plan()
        return {
            'model': model,
            'buyer': buyer,
            'beneficiary': beneficiary,
            'production': production,
            'lines': lines,
            'products': products,
            'boms': boms,
            'auto_fabric': auto_fabric,
            'manual_fabric': manual_fabric,
            'second_fabric': second_fabric,
            'auto_takawe': auto_takawe,
            'manual_takawe': manual_takawe,
            'other_material': other_material,
        }

    def _open_setup(self, production, user=False):
        target = production.with_user(user) if user else production
        action = target.action_open_tailoring_material_setup()
        self.assertEqual(action['type'], 'ir.actions.act_window')
        self.assertEqual(
            action['res_model'],
            'furniture.mrp.tailoring.setup.wizard',
        )
        self.assertEqual(action['target'], 'new')
        self.assertEqual(
            action['view_id'],
            self.env.ref(
                'furniture_mrp.view_furniture_mrp_tailoring_setup_wizard_form'
            ).id,
        )
        wizard = self.env[action['res_model']].browse(
            action['res_id'],
        ).exists()
        self.assertTrue(wizard)
        self.assertEqual(wizard.production_id, production)
        return wizard

    def _open_editor(self, production_line, kind):
        setup = self._open_setup(production_line.production_id)
        setup_line = setup.line_ids.filtered(
            lambda row: row.production_line_id == production_line
        ).ensure_one()
        method_name = (
            'action_open_fabric_editor'
            if kind == 'fabric'
            else 'action_open_takawe_editor'
        )
        action = getattr(setup_line, method_name)()
        self.assertEqual(action['type'], 'ir.actions.client')
        self.assertEqual(
            action['tag'],
            'furniture_tailoring_material_popup_open',
        )
        params = action['params']
        self.assertEqual(
            params['res_model'],
            'furniture.mrp.tailoring.material.wizard',
        )
        self.assertEqual(
            params['view_id'],
            self.env.ref(
                'furniture_mrp.view_furniture_mrp_tailoring_material_wizard_form'
            ).id,
        )
        self.assertEqual(params['size'], 'xl')
        self.assertEqual(
            params['context']['form_view_initial_mode'],
            'edit',
        )
        editor = self.env[params['res_model']].browse(
            params['res_id'],
        ).exists()
        self.assertTrue(editor)
        self.assertEqual(editor.production_line_id, production_line)
        self.assertEqual(editor.material_kind, kind)
        return editor

    def _apply_editor(self, production_line, kind, material_values):
        editor = self._open_editor(production_line, kind)
        editor.write({
            'line_ids': [(5, 0, 0)] + [
                (0, 0, {
                    'product_id': product.id,
                    'qty': qty,
                    'product_uom_id': (uom or product.uom_id).id,
                    'piece_size': '45' if kind == 'takawe' else False,
                })
                for product, qty, uom in material_values
            ],
        })
        editor.action_apply()
        production_line.invalidate_recordset([
            'tailoring_fabric_configured',
            'tailoring_takawe_configured',
        ])
        return editor

    def _stage_materials(self, production_line, stage):
        return production_line.production_id.material_line_ids.filtered(
            lambda material: (
                material.production_line_id == production_line
                and material.stage == stage
            )
        )

    def _piece_image_attachments(self, production_line):
        return self.env['ir.attachment'].sudo().search([
            ('res_model', '=', 'furniture.mrp.production.line'),
            ('res_id', '=', production_line.id),
            ('res_field', '=', 'batch_image_1920'),
        ])

    def _order_image_attachments(self, production):
        return self.env['ir.attachment'].sudo().search([
            ('res_model', '=', 'furniture.mrp.production'),
            ('res_id', '=', production.id),
            ('res_field', '=', 'tailoring_set_image_1920'),
        ])

    def _tailoring_materials(self, production_line):
        return self._stage_materials(production_line, 'tailoring')

    def _upholstery_materials(self, production_line):
        return self._stage_materials(production_line, 'upholstery')

    def _allocation_kind_field(self):
        Allocation = self.env[
            'furniture.mrp.tailoring.material.allocation'
        ]
        for field_name in ('material_kind', 'kind'):
            if field_name in Allocation._fields:
                return field_name
        self.fail('The tailoring material allocation has no kind field.')

    def _allocation_relation_field(self):
        matches = [
            field_name
            for field_name, field in self.env[
                'furniture.mrp.production.line'
            ]._fields.items()
            if (
                field.type == 'one2many'
                and field.comodel_name
                == 'furniture.mrp.tailoring.material.allocation'
            )
        ]
        self.assertTrue(
            matches,
            'Production lines need an explicit allocation relation.',
        )
        return (
            'tailoring_material_allocation_ids'
            if 'tailoring_material_allocation_ids' in matches
            else matches[0]
        )

    def _allocation_values(
        self, production_line, kind, product, qty, uom=False,
        production=False,
    ):
        Allocation = self.env[
            'furniture.mrp.tailoring.material.allocation'
        ]
        values = {
            'production_line_id': production_line.id,
            self._allocation_kind_field(): kind,
            'product_id': product.id,
            'qty': qty,
            'product_uom_id': (uom or product.uom_id).id,
            'piece_size': '45' if kind == 'takawe' else False,
        }
        if 'production_id' in Allocation._fields:
            values['production_id'] = (
                production or production_line.production_id
            ).id
        return values

    def _create_allocation(
        self, production_line, kind, product, qty, uom=False,
        production=False,
    ):
        return self.env[
            'furniture.mrp.tailoring.material.allocation'
        ].create(self._allocation_values(
            production_line,
            kind,
            product,
            qty,
            uom=uom,
            production=production,
        ))

    def _create_store_request(self, fixture, state):
        production = fixture['production']
        return self.env['furniture.mrp.store.request'].sudo().create({
            'production_id': production.id,
            'stage_code': 'tailoring',
            'stage_order_model': 'furniture.mrp.tailoring',
            'stage_order_res_id': 900000 + production.id,
            'stage_order_name': 'TAL/SETUP/LOCK',
            'request_kind': 'direct',
            'start_mode': 'direct',
            'state': state,
            'requested_by_id': self.env.user.id,
            'assigned_to_id': self.env.user.id,
            'requested_product_summary': 'Tailoring setup lock test',
            'payload_json': {
                'production_line_ids': fixture['lines'].ids,
            },
        })

    def test_action_builds_three_compact_rows_with_product_context(self):
        fixture = self._create_fixture()

        wizard = self._open_setup(fixture['production'])

        self.assertEqual(len(wizard.line_ids), 3)
        for row, production_line in zip(
            wizard.line_ids,
            fixture['lines'],
        ):
            self.assertEqual(row.production_line_id, production_line)
            self.assertEqual(row.product_id, production_line.product_id)
            self.assertEqual(row.model_id, fixture['model'])
            self.assertEqual(row.product_qty, production_line.product_qty)
            self.assertEqual(
                row.product_uom_id,
                production_line.product_uom_id,
            )
        self.assertEqual(wizard.buyer_partner_id, fixture['buyer'])
        self.assertEqual(
            wizard.beneficiary_partner_id,
            fixture['beneficiary'],
        )
        self.assertTrue(all(row.product_id.display_name for row in wizard.line_ids))
        self.assertTrue(all(row.model_id.display_name for row in wizard.line_ids))

    def test_material_kind_is_shared_by_template_and_variant(self):
        fixture = self._create_fixture()
        template_field = self.env[
            'product.template'
        ]._fields['furniture_tailoring_material_kind']
        variant_field = self.env[
            'product.product'
        ]._fields['furniture_tailoring_material_kind']

        self.assertFalse(template_field.required)
        self.assertFalse(variant_field.readonly)
        self.assertEqual(
            template_field.get_values(self.env),
            ['fabric', 'takawe'],
        )

        for product, expected_kind in (
            (fixture['manual_fabric'], 'fabric'),
            (fixture['manual_takawe'], 'takawe'),
        ):
            self.assertIn(
                'furniture_tailoring_material_kind',
                product.product_tmpl_id._fields,
            )
            self.assertIn(
                'furniture_tailoring_material_kind',
                product._fields,
            )
            self.assertEqual(
                product.product_tmpl_id.furniture_tailoring_material_kind,
                expected_kind,
            )
            self.assertEqual(
                product.furniture_tailoring_material_kind,
                expected_kind,
            )

        neutral = self._new_raw_product('قماش بالاسم فقط')
        self.assertFalse(neutral.furniture_tailoring_material_kind)
        self.assertFalse(neutral._furniture_tailoring_material_kind())
        neutral.write({'furniture_tailoring_material_kind': 'fabric'})
        self.assertEqual(
            neutral.product_tmpl_id.furniture_tailoring_material_kind,
            'fabric',
        )
        self.assertEqual(
            self.env['product.product'].search_count([
                ('id', '=', neutral.id),
                ('furniture_tailoring_material_kind', '=', 'fabric'),
            ]),
            1,
        )
        neutral.write({'furniture_tailoring_material_kind': False})
        self.assertFalse(
            neutral.product_tmpl_id.furniture_tailoring_material_kind
        )
        self.assertFalse(neutral._furniture_tailoring_material_kind())

    def test_material_kind_is_exposed_in_inventory_product_and_stock_views(self):
        view_expectations = {
            'furniture_mrp.'
            'view_product_product_tree_furniture_display_name': 1,
            'furniture_mrp.'
            'view_product_template_tree_furniture_display_name': 1,
            'furniture_mrp.view_stock_quant_tree_furniture_model': 1,
            'furniture_mrp.view_stock_quant_tree_simple_furniture_model': 1,
            'furniture_mrp.view_stock_quant_inventory_furniture_model': 1,
            'furniture_mrp.view_product_template_form_furniture_supplier': 1,
        }
        for xml_id, expected_count in view_expectations.items():
            view = self.env.ref(xml_id)
            view._check_xml()
            document = etree.fromstring(view.arch_db.encode())
            fields_found = document.xpath(
                ".//field[@name='furniture_tailoring_material_kind']"
            )
            self.assertEqual(
                len(fields_found),
                expected_count,
                xml_id,
            )

        quant_field = self.env['stock.quant']._fields[
            'furniture_tailoring_material_kind'
        ]
        self.assertTrue(quant_field.readonly)
        self.assertEqual(
            quant_field.related,
            'product_id.furniture_tailoring_material_kind',
        )

        for model_name, view_type in (
            ('product.template', 'list'),
            ('product.template', 'form'),
            ('product.product', 'list'),
            ('product.product', 'form'),
            ('stock.quant', 'list'),
        ):
            combined_arch = self.env[model_name].get_view(
                view_type=view_type,
            )['arch']
            combined_document = etree.fromstring(combined_arch.encode())
            self.assertEqual(
                len(combined_document.xpath(
                    ".//field["
                    "@name='furniture_tailoring_material_kind'"
                    "]"
                )),
                1,
                '%s %s' % (model_name, view_type),
            )

    def test_main_view_uses_product_specific_material_popups(self):
        view = self.env.ref(
            'furniture_mrp.view_furniture_mrp_tailoring_setup_wizard_form'
        )
        view._check_xml()
        document = etree.fromstring(view.arch_db.encode())
        self.assertEqual(document.get('edit'), '1')
        self.assertEqual(
            document.get('js_class'),
            'furniture_tailoring_setup_inline_form',
        )
        order_identity = document.xpath(
            "//div[contains(concat(' ', normalize-space(@class), ' '), "
            "' o_furniture_tailoring_order_identity ')]"
        )
        self.assertEqual(len(order_identity), 1)
        self.assertEqual(
            ''.join(order_identity[0].xpath('./span/text()')).strip(),
            'الموديل',
        )
        self.assertEqual(
            order_identity[0].xpath('./strong/field/@name'),
            ['model_id'],
        )
        first_meta_chip = document.xpath(
            "//div[contains(concat(' ', normalize-space(@class), ' '), "
            "' o_furniture_tailoring_setup_meta ')]/div[1]"
        )
        self.assertEqual(len(first_meta_chip), 1)
        self.assertEqual(
            ''.join(first_meta_chip[0].xpath('./span/text()')).strip(),
            'أمر الإنتاج',
        )
        self.assertEqual(
            first_meta_chip[0].xpath('./strong/field/@name'),
            ['order_name'],
        )
        line_fields = document.xpath("//field[@name='line_ids']")
        self.assertEqual(len(line_fields), 1)
        line_field = line_fields[0]
        self.assertEqual(line_field.get('mode'), 'kanban')
        self.assertEqual(len(line_field.xpath('./kanban')), 1)
        self.assertFalse(line_field.xpath('./list'))
        self.assertFalse(document.xpath(
            "//section[contains(concat(' ', normalize-space(@class), ' '), "
            "' o_furniture_tailoring_inline_editor ')]"
        ))
        self.assertFalse(document.xpath("//field[@name='editor_line_ids']"))

        kanban = line_field.xpath('./kanban')[0]
        self.assertEqual(kanban.get('create'), '0')
        self.assertEqual(kanban.get('can_open'), '0')
        card_articles = kanban.xpath(".//t[@t-name='card']/article")
        self.assertEqual(len(card_articles), 1)
        self.assertEqual(
            [child.get('class') for child in card_articles[0]],
            [
                'o_furniture_tailoring_setup_notes',
                'o_furniture_tailoring_setup_material_cell is-takawe',
                'o_furniture_tailoring_setup_material_cell is-fabric',
                'o_furniture_tailoring_setup_product',
            ],
        )
        expected_actions = {
            'action_open_fabric_editor',
            'action_open_takawe_editor',
        }
        plus_buttons = kanban.xpath(
            ".//button[@type='object' and "
            "(@name='action_open_fabric_editor' or "
            "@name='action_open_takawe_editor')]"
        )
        self.assertEqual(
            {button.get('name') for button in plus_buttons},
            expected_actions,
        )
        self.assertTrue(all(
            '+' in button.get('string', '') for button in plus_buttons
        ))
        takawe_titles = kanban.xpath(
            ".//div[contains(concat(' ', normalize-space(@class), ' '), "
            "' is-takawe ')]//span["
            "contains(concat(' ', normalize-space(@class), ' '), "
            "' o_furniture_tailoring_material_button_title ')"
            "]"
        )
        self.assertEqual(len(takawe_titles), 2)
        self.assertTrue(all(
            title.get('t-if') == '!record.takawe_summary.raw_value'
            for title in takawe_titles
        ))
        fabric_titles = kanban.xpath(
            ".//div[contains(concat(' ', normalize-space(@class), ' '), "
            "' is-fabric ')]//span["
            "contains(concat(' ', normalize-space(@class), ' '), "
            "' o_furniture_tailoring_material_button_title ')"
            "]"
        )
        self.assertEqual(len(fabric_titles), 2)
        self.assertTrue(all(
            title.get('t-if') == '!record.fabric_summary.raw_value'
            for title in fabric_titles
        ))
        self.assertFalse(kanban.xpath(
            ".//button[@name='action_select_fabric_editor' or "
            "@name='action_select_takawe_editor' or "
            "@name='action_open_image_preview' or "
            "@name='action_select_image_editor']"
        ))
        direct_image_fields = kanban.xpath(
            ".//field[@name='image_128' and "
            "@widget='furniture_piece_image_upload']"
        )
        self.assertFalse(direct_image_fields)
        self.assertFalse(kanban.xpath(".//field[@name='image_128']"))
        order_image_fields = document.xpath(
            "//div[contains(@class, 'o_furniture_tailoring_order_image')]/"
            "field[@name='order_image_128']"
        )
        self.assertEqual(len(order_image_fields), 2)
        order_uploaders = [
            field for field in order_image_fields
            if field.get('widget') == 'furniture_order_image_upload'
        ]
        order_previews = [
            field for field in order_image_fields
            if field.get('widget') == 'furniture_order_image_preview'
        ]
        self.assertEqual(len(order_uploaders), 1)
        self.assertEqual(len(order_previews), 1)
        accepted_types = order_uploaders[0].get('options', '')
        for mimetype in ('image/jpeg', 'image/png'):
            self.assertIn(mimetype, accepted_types)
        self.assertFalse(document.xpath(
            "//div[contains(@class, 'o_furniture_tailoring_setup_intro')]"
        ))
        note_inputs = kanban.xpath(
            ".//field[@name='piece_note' and "
            "@widget='furniture_piece_note_input']"
        )
        self.assertEqual(len(note_inputs), 1)
        editor_view = self.env.ref(
            'furniture_mrp.'
            'view_furniture_mrp_tailoring_material_wizard_form'
        )
        editor_view._check_xml()
        editor_document = etree.fromstring(
            editor_view.arch_db.encode()
        )
        material_selector_fields = editor_document.xpath(
            ".//field[@name='selected_material_product_ids' and "
            "@widget='furniture_multi_card_selector']"
        )
        self.assertEqual(len(material_selector_fields), 2)
        for material_selector in material_selector_fields:
            strict_domain = material_selector.get('domain', '')
            self.assertIn(
                "('furniture_tailoring_material_kind', 'in', "
                "['fabric', 'takawe'])",
                strict_domain,
            )
            self.assertNotIn(
                "('furniture_tailoring_material_kind', '=', False)",
                strict_domain,
            )
            selector_options = material_selector.get('options', '')
            self.assertIn("'quick_limit': 8", selector_options)
            self.assertIn("'card_variant': 'material'", selector_options)
        fabric_selectors = [
            field for field in material_selector_fields
            if "material_kind != 'fabric'" in field.get('invisible', '')
        ]
        takawe_selectors = [
            field for field in material_selector_fields
            if "material_kind == 'fabric'" in field.get('invisible', '')
        ]
        self.assertEqual(len(fabric_selectors), 1)
        self.assertEqual(len(takawe_selectors), 1)
        self.assertIn(
            "'single_selection': True",
            fabric_selectors[0].get('options', ''),
        )
        self.assertNotIn(
            "'single_selection'",
            takawe_selectors[0].get('options', ''),
        )
        editor_lists = editor_document.xpath("//field[@name='line_ids']/list")
        self.assertEqual(len(editor_lists), 1)
        self.assertEqual(editor_lists[0].get('create'), '0')
        self.assertEqual(editor_lists[0].get('delete'), '0')
        self.assertFalse(editor_lists[0].xpath('./control/create'))
        readonly_material_fields = editor_lists[0].xpath(
            "./field[@name='product_id' and @readonly='1']"
        )
        self.assertEqual(len(readonly_material_fields), 1)
        editor_uom_fields = editor_document.xpath(
            "//field[@name='line_ids']/list/"
            "field[@name='product_uom_id']"
        )
        self.assertEqual(len(editor_uom_fields), 1)
        self.assertEqual(editor_uom_fields[0].get('readonly'), '1')
        editor_qty_fields = editor_document.xpath(
            "//field[@name='line_ids']/list/field[@name='qty']"
        )
        self.assertEqual(len(editor_qty_fields), 1)
        self.assertEqual(
            editor_qty_fields[0].get('readonly'),
            "parent.material_kind == 'fabric'",
        )
        size_fields = editor_document.xpath(
            "//field[@name='line_ids']/list/"
            "field[@name='piece_size' and "
            "@widget='furniture_touch_selection']"
        )
        self.assertEqual(len(size_fields), 1)
        self.assertEqual(
            size_fields[0].get('required'),
            "parent.material_kind == 'takawe'",
        )
        self.assertEqual(
            size_fields[0].get('column_invisible'),
            "parent.material_kind != 'takawe'",
        )
        size_sections = editor_document.xpath(
            "//div[contains(@class, "
            "'o_furniture_tailoring_piece_size_section')]"
        )
        self.assertFalse(size_sections)
        quick_selector_js = (
            Path(__file__).resolve().parents[1]
            / 'static/src/js/furniture_quick_card_selector.js'
        ).read_text()
        self.assertIn('class FurnitureTouchSelection', quick_selector_js)
        self.assertIn(
            '.add("furniture_touch_selection", {',
            quick_selector_js,
        )
        self.assertIn(
            'await this.props.record.update({ [this.props.name]: value });',
            quick_selector_js,
        )
        quick_selector_templates = etree.parse(str(
            Path(__file__).resolve().parents[1]
            / 'static/src/xml/furniture_quick_card_selector.xml'
        ))
        touch_template = quick_selector_templates.xpath(
            "//t[@t-name='furniture_mrp.FurnitureTouchSelection']"
        )
        self.assertEqual(len(touch_template), 1)
        touch_buttons = touch_template[0].xpath(
            ".//button[@role='radio' and @type='button']"
        )
        self.assertEqual(len(touch_buttons), 1)
        self.assertEqual(
            touch_buttons[0].get('t-on-click.stop.prevent'),
            '() => this.selectValue(option[0])',
        )
        self.assertFalse(touch_template[0].xpath(".//i"))
        touch_css = (
            Path(__file__).resolve().parents[1]
            / 'static/src/css/furniture_mrp.css'
        ).read_text()
        self.assertIn('.o_furniture_touch_selection__button', touch_css)
        self.assertIn(
            'grid-template-columns: repeat(3, 52px) 78px;',
            touch_css,
        )
        self.assertIn('height: 34px;', touch_css)
        self.assertIn('min-height: 40px;', touch_css)
        self.assertIn('white-space: nowrap;', touch_css)
        self.assertEqual(
            len(editor_document.xpath(
                "//div[contains(@class, "
                "'o_furniture_tailoring_material_editor_context')]/"
                "span[contains(@class, 'is-product')]"
            )),
            1,
        )
        self.assertEqual(
            len(editor_document.xpath(
                "//div[contains(@class, "
                "'o_furniture_tailoring_material_editor_context')]/"
                "span[contains(@class, 'is-model')]"
            )),
            1,
        )
        self.assertFalse(kanban.xpath(
            ".//field[@name='fabric_piece_size']"
        ))
        self.assertFalse(kanban.xpath(
            ".//span[contains(@class, 'is-fabric')]"
        ))
        self.assertFalse(kanban.xpath(
            ".//field[@name='takawe_piece_size']"
        ))
        self.assertFalse(kanban.xpath(
            ".//div[contains(@class, "
            "'o_furniture_tailoring_setup_material_cell')]//"
            "span[contains(@class, "
            "'o_furniture_tailoring_material_size')]"
        ))
        apply_buttons = editor_document.xpath(
            "//button[@name='action_apply' and @type='object']"
        )
        self.assertEqual(len(apply_buttons), 1)
        self.assertIsNone(apply_buttons[0].get('close'))
        cancel_buttons = editor_document.xpath(
            "//button[@special='cancel' and @close='1' and @string='إلغاء']"
        )
        self.assertEqual(len(cancel_buttons), 1)
        self.assertFalse(editor_document.xpath(
            "//button[@name='action_back_to_setup' and @string='إلغاء']"
        ))
        self.assertFalse(editor_document.xpath(
            "//div[contains(@class, "
            "'o_furniture_tailoring_material_editor_icon')]"
        ))
        self.assertFalse(editor_document.xpath(
            "//div[contains(@class, "
            "'o_furniture_tailoring_material_editor_header')]//"
            "i[contains(@class, 'fa-plus')]"
        ))

        popup_js = (
            Path(__file__).resolve().parents[1]
            / 'static/src/js/tailoring_setup_inline_form.js'
        ).read_text()
        self.assertIn('request.resolve(true);', popup_js)
        self.assertIn('resolve(false);', popup_js)
        self.assertIn(
            'import { FormViewDialog } from '
            '"@web/views/view_dialogs/form_view_dialog";',
            popup_js,
        )
        self.assertIn('closeDialog = env.services.dialog.add(', popup_js)
        self.assertIn('const tailoringMaterialPopupClosers = new Map();', popup_js)
        self.assertIn(
            'tailoringMaterialPopupClosers.set(popupKey, closeDialog);',
            popup_js,
        )
        self.assertIn(
            'tailoringMaterialPopupClosers.get(popupKey)',
            popup_js,
        )
        self.assertIn(
            '.add("furniture_tailoring_material_popup_open", '
            'openTailoringMaterialPopup)',
            popup_js,
        )
        self.assertIn(
            'await refreshTailoringSetupInlineForm(env, action);',
            popup_js,
        )
        for gallery_fragment in (
            'export class FurnitureOrderImageGallery',
            'get_order_image_gallery',
            'add_order_image',
            'remove_order_image',
            'o_furniture_tailoring_order_gallery_add',
            'fa fa-plus',
        ):
            self.assertIn(gallery_fragment, popup_js)
        popup_css = (
            Path(__file__).resolve().parents[1]
            / 'static/src/css/furniture_mrp.css'
        ).read_text()
        self.assertIn(
            '.o_furniture_tailoring_order_gallery_add',
            popup_css,
        )
        self.assertNotIn('action.params?.setup_action', popup_js)
        self.assertNotIn('type: "ir.actions.act_window_close"', popup_js)
        self.assertEqual(
            len(document.xpath("//field[@name='piece_image_1920']")),
            0,
        )

        production_view = self.env.ref(
            'furniture_mrp.view_furniture_mrp_production_form'
        )
        admin = self.env.ref('base.user_admin')
        effective_arch = self.env[production_view.model].with_user(
            admin
        ).get_view(
            view_id=production_view.id,
            view_type='form',
        )['arch']
        effective_document = etree.fromstring(effective_arch.encode())
        self.assertEqual(
            len(effective_document.xpath(
                "//button[@name='action_open_kit_planner']"
            )),
            0,
            'The production form must not expose the removed Kit planner button.',
        )
        production_document = etree.fromstring(
            production_view.arch_db.encode()
        )
        production_buttons = production_document.xpath(
            "//div[contains(concat(' ', normalize-space(@class), ' '), "
            "' o_furniture_products_workbench_actions ')]/button"
            "[@name='action_open_kit_planner']"
        )
        self.assertFalse(production_buttons)
        material_buttons = production_document.xpath(
            "//div[contains(concat(' ', normalize-space(@class), ' '), "
            "' o_furniture_products_workbench_actions ')]/button"
            "[@name='action_open_tailoring_material_setup']"
        )
        self.assertEqual(len(material_buttons), 1)
        self.assertEqual(material_buttons[0].get('type'), 'object')
        self.assertEqual(
            material_buttons[0].get('string'),
            'الأقمشة والتكاوي',
        )
        self.assertIn(
            'furniture_mrp.group_furniture_mrp_manager',
            material_buttons[0].get('groups', ''),
        )
        compact_view = self.env.ref(
            'furniture_mrp.view_furniture_mrp_tailoring_setup_wizard_form'
        )
        compact_document = etree.fromstring(compact_view.arch_db.encode())
        self.assertFalse(
            compact_document.xpath("//field[@name='kit_display_label']")
        )
        self.assertEqual(
            len(compact_document.xpath(
                "//div[contains(@class, "
                "'o_furniture_tailoring_setup_chip')]"
            )),
            3,
        )

    def test_material_cards_open_product_specific_popups(self):
        fixture = self._create_fixture()
        setup = self._open_setup(fixture['production'])
        setup_line = setup.line_ids.filtered(
            lambda row: row.production_line_id == fixture['lines'][0]
        ).ensure_one()

        for method_name, expected_kind in (
            ('action_open_fabric_editor', 'fabric'),
            ('action_open_takawe_editor', 'takawe'),
        ):
            action = getattr(setup_line, method_name)()
            self.assertEqual(action['type'], 'ir.actions.client')
            self.assertEqual(
                action['tag'],
                'furniture_tailoring_material_popup_open',
            )
            params = action['params']
            self.assertEqual(
                params['res_model'],
                'furniture.mrp.tailoring.material.wizard',
            )
            editor = self.env[params['res_model']].browse(
                params['res_id']
            ).exists()
            self.assertTrue(editor)
            self.assertEqual(editor.setup_wizard_id, setup)
            self.assertEqual(
                editor.production_line_id,
                fixture['lines'][0],
            )
            self.assertEqual(editor.material_kind, expected_kind)
            self.assertEqual(
                set(editor.selected_material_product_ids.ids),
                set(editor.line_ids.mapped('product_id').ids),
            )

    def test_fabric_card_change_keeps_recipe_quantity(self):
        fixture = self._create_fixture()
        editor = self._open_editor(fixture['lines'][0], 'fabric')
        meter = self.env.ref('furniture_mrp.furniture_uom_meter')

        editor.selected_material_product_ids = fixture['manual_fabric']
        editor._onchange_selected_material_product_ids()
        manual_line = editor.line_ids.ensure_one()
        self.assertEqual(manual_line.product_id, fixture['manual_fabric'])
        self.assertAlmostEqual(manual_line.qty, 4.0)
        self.assertEqual(editor.line_ids.mapped('product_uom_id'), meter)

        manual_line.qty = 99.0
        editor.selected_material_product_ids = fixture['second_fabric']
        editor._onchange_selected_material_product_ids()
        replacement_line = editor.line_ids.ensure_one()
        self.assertEqual(replacement_line.product_id, fixture['second_fabric'])
        self.assertAlmostEqual(replacement_line.qty, 4.0)

    def test_takawe_cards_keep_multiple_editable_quantities(self):
        fixture = self._create_fixture()
        editor = self._open_editor(fixture['lines'][0], 'takawe')
        editor.selected_material_product_ids = (
            fixture['auto_takawe'] | fixture['manual_takawe']
        )
        editor._onchange_selected_material_product_ids()
        self.assertEqual(len(editor.line_ids), 2)
        manual_line = editor.line_ids.filtered(
            lambda line: line.product_id == fixture['manual_takawe']
        ).ensure_one()
        self.assertAlmostEqual(manual_line.qty, 1.0)
        manual_line.qty = 2.75
        manual_line.piece_size = '55'
        auto_line = editor.line_ids.filtered(
            lambda line: line.product_id == fixture['auto_takawe']
        ).ensure_one()
        auto_line.piece_size = '40x60'
        editor._onchange_selected_material_product_ids()
        preserved_line = editor.line_ids.filtered(
            lambda line: line.product_id == fixture['manual_takawe']
        ).ensure_one()
        self.assertAlmostEqual(preserved_line.qty, 2.75)
        self.assertEqual(preserved_line.piece_size, '55')
        self.assertEqual(
            editor.line_ids.filtered(
                lambda line: line.product_id == fixture['auto_takawe']
            ).piece_size,
            '40x60',
        )

        editor.action_apply()
        allocations = fixture['lines'][0].tailoring_material_allocation_ids.filtered(
            lambda allocation: allocation.material_kind == 'takawe'
        )
        self.assertEqual(len(allocations), 2)
        self.assertEqual(
            {
                allocation.product_id: allocation.piece_size
                for allocation in allocations
            },
            {
                fixture['auto_takawe']: '40x60',
                fixture['manual_takawe']: '55',
            },
        )
        fixture['lines'][0].invalidate_recordset([
            'tailoring_takawe_piece_size',
        ])
        self.assertFalse(fixture['lines'][0].tailoring_takawe_piece_size)

    def test_only_takawe_popup_requires_exact_piece_size_choice(self):
        size_field = self.env[
            'furniture.mrp.tailoring.material.wizard.line'
        ]._fields['piece_size']
        self.assertEqual(
            size_field.selection,
            [
                ('45', '45'),
                ('50', '50'),
                ('55', '55'),
                ('40x60', '40×60'),
            ],
        )

        fixture = self._create_fixture()
        fabric_editor = self._open_editor(fixture['lines'][0], 'fabric')
        self.assertFalse(fabric_editor.piece_size)
        fabric_editor.action_apply()
        fixture['lines'][0].invalidate_recordset([
            'tailoring_fabric_piece_size',
        ])
        self.assertFalse(fixture['lines'][0].tailoring_fabric_piece_size)

        editor = self._open_editor(fixture['lines'][0], 'takawe')
        self.assertFalse(editor.line_ids.ensure_one().piece_size)
        with self.assertRaisesRegex(ValidationError, 'اختار المقاس'):
            editor.action_apply()

    def test_recipe_fabric_is_fixed_and_takawe_is_editable_default(self):
        fixture = self._create_fixture()
        production_line = fixture['lines'][0]
        meter = self.env.ref('furniture_mrp.furniture_uom_meter')
        setup = self._open_setup(fixture['production'])
        setup_line = setup.line_ids.filtered(
            lambda row: row.production_line_id == production_line
        ).ensure_one()

        self.assertIn(fixture['auto_fabric'].display_name, setup_line.fabric_summary)
        self.assertIn(fixture['auto_takawe'].display_name, setup_line.takawe_summary)
        self.assertFalse(production_line.tailoring_material_allocation_ids)
        self.assertFalse(production_line.tailoring_fabric_configured)
        self.assertFalse(production_line.tailoring_takawe_configured)

        fabric_action = setup_line.action_open_fabric_editor()
        fabric_params = fabric_action['params']
        fabric_editor = self.env[fabric_params['res_model']].browse(
            fabric_params['res_id']
        )
        fabric_draft = fabric_editor.line_ids.ensure_one()
        self.assertEqual(fabric_draft.product_id, fixture['auto_fabric'])
        self.assertAlmostEqual(fabric_draft.qty, 4.0)
        self.assertEqual(fabric_draft.product_uom_id, meter)
        self.assertFalse(fabric_editor.piece_size)

        takawe_action = setup_line.action_open_takawe_editor()
        takawe_params = takawe_action['params']
        takawe_editor = self.env[takawe_params['res_model']].browse(
            takawe_params['res_id']
        )
        takawe_draft = takawe_editor.line_ids.ensure_one()
        self.assertEqual(takawe_draft.product_id, fixture['auto_takawe'])
        self.assertAlmostEqual(takawe_draft.qty, 2.0)
        self.assertEqual(takawe_draft.product_uom_id, meter)

        takawe_editor.write({
            'line_ids': [(5, 0, 0), (0, 0, {
                'product_id': fixture['manual_takawe'].id,
                'qty': 3.5,
                'product_uom_id': fixture['manual_takawe'].uom_id.id,
                'piece_size': '40x60',
            })],
        })
        save_action = takawe_editor.action_apply()
        self.assertEqual(save_action['type'], 'ir.actions.client')
        self.assertEqual(
            save_action['tag'],
            'furniture_tailoring_material_popup_saved',
        )
        self.assertEqual(save_action['params']['editor_id'], takawe_editor.id)
        self.assertEqual(save_action['params']['wizard_id'], setup.id)
        self.assertEqual(
            save_action['params']['res_model'],
            'furniture.mrp.tailoring.setup.wizard',
        )
        self.assertEqual(
            save_action['params']['res_id'],
            setup.id,
        )
        production_line.invalidate_recordset([
            'tailoring_material_allocation_ids',
            'tailoring_takawe_configured',
            'tailoring_fabric_piece_size',
            'tailoring_takawe_piece_size',
        ])
        self.assertTrue(production_line.tailoring_takawe_configured)
        self.assertFalse(production_line.tailoring_fabric_piece_size)
        self.assertEqual(
            production_line.tailoring_takawe_piece_size,
            '40x60',
        )
        saved_takawe = (
            production_line.tailoring_material_allocation_ids.filtered(
                lambda allocation: allocation.material_kind == 'takawe'
            ).ensure_one()
        )
        self.assertEqual(saved_takawe.product_id, fixture['manual_takawe'])
        self.assertAlmostEqual(saved_takawe.qty, 3.5)
        self.assertEqual(saved_takawe.product_uom_id, meter)
        self.assertEqual(saved_takawe.piece_size, '40x60')

        reopened_setup = self._open_setup(fixture['production'])
        reopened_row = reopened_setup.line_ids.filtered(
            lambda row: row.production_line_id == production_line
        ).ensure_one()
        self.assertFalse(reopened_row.fabric_piece_size)
        self.assertEqual(reopened_row.takawe_piece_size, '40x60')
        reopened_action = reopened_row.action_open_takawe_editor()
        reopened_params = reopened_action['params']
        reopened_editor = self.env[reopened_params['res_model']].browse(
            reopened_params['res_id']
        )
        reopened_draft = reopened_editor.line_ids.ensure_one()
        self.assertEqual(reopened_draft.product_id, fixture['manual_takawe'])
        self.assertAlmostEqual(reopened_draft.qty, 3.5)
        self.assertEqual(reopened_draft.product_uom_id, meter)
        self.assertEqual(reopened_draft.piece_size, '40x60')

    def test_fabric_apply_ignores_forged_qty_and_rejects_multiple_products(self):
        fixture = self._create_fixture()
        line = fixture['lines'][0]
        editor = self._open_editor(line, 'fabric')
        editor.write({
            'piece_size': '50',
            'line_ids': [(5, 0, 0), (0, 0, {
                'product_id': fixture['manual_fabric'].id,
                'qty': 99.0,
                'product_uom_id': fixture['manual_fabric'].uom_id.id,
            })],
        })
        editor.action_apply()
        allocation = line.tailoring_material_allocation_ids.filtered(
            lambda item: item.material_kind == 'fabric'
        ).ensure_one()
        self.assertEqual(allocation.product_id, fixture['manual_fabric'])
        self.assertAlmostEqual(allocation.qty, 4.0)
        line.invalidate_recordset(['tailoring_fabric_piece_size'])
        self.assertFalse(line.tailoring_fabric_piece_size)

        reopened = self._open_editor(line, 'fabric')
        reopened.write({
            'line_ids': [(5, 0, 0), (0, 0, {
                'product_id': fixture['manual_fabric'].id,
                'qty': 4.0,
                'product_uom_id': fixture['manual_fabric'].uom_id.id,
            }), (0, 0, {
                'product_id': fixture['second_fabric'].id,
                'qty': 4.0,
                'product_uom_id': fixture['second_fabric'].uom_id.id,
            })],
        })
        with self.assertRaisesRegex(ValidationError, 'نوع قماش واحد'):
            reopened.action_apply()

    def test_only_takawe_keeps_piece_size(self):
        fixture = self._create_fixture()
        production_line = fixture['lines'][0]

        production_line.sudo().with_context(
            furniture_tailoring_setup_internal_write=True,
        ).write({'tailoring_fabric_piece_size': '50'})
        production_line.invalidate_recordset([
            'tailoring_fabric_piece_size',
        ])
        self.assertFalse(production_line.tailoring_fabric_piece_size)

        fabric_editor = self._open_editor(production_line, 'fabric')
        fabric_editor.action_apply()

        takawe_editor = self._open_editor(production_line, 'takawe')
        self.assertFalse(takawe_editor.line_ids.ensure_one().piece_size)
        takawe_editor.line_ids.ensure_one().piece_size = '55'
        takawe_editor.action_apply()

        production_line.invalidate_recordset([
            'tailoring_fabric_piece_size',
            'tailoring_takawe_piece_size',
        ])
        self.assertFalse(production_line.tailoring_fabric_piece_size)
        self.assertEqual(production_line.tailoring_takawe_piece_size, '55')

        reopened_setup = self._open_setup(fixture['production'])
        reopened_row = reopened_setup.line_ids.filtered(
            lambda row: row.production_line_id == production_line
        ).ensure_one()
        self.assertFalse(reopened_row.fabric_piece_size)
        self.assertEqual(reopened_row.takawe_piece_size, '55')

        reopened_fabric = self._open_editor(production_line, 'fabric')
        reopened_takawe = self._open_editor(production_line, 'takawe')
        self.assertFalse(reopened_fabric.line_ids.ensure_one().piece_size)
        self.assertEqual(
            reopened_takawe.line_ids.ensure_one().piece_size,
            '55',
        )

    def test_inline_cancel_discards_draft_without_persistent_mutation(self):
        fixture = self._create_fixture()
        production = fixture['production']
        production_line = fixture['lines'][0]
        setup = self._open_setup(production)
        setup_line = setup.line_ids.filtered(
            lambda row: row.production_line_id == production_line
        ).ensure_one()
        baseline_revision = production.tailoring_material_revision
        baseline_allocation_ids = (
            production_line.tailoring_material_allocation_ids.ids
        )
        baseline_materials = [
            (
                material.id,
                material.product_id.id,
                material.stage,
                material.qty_needed,
            )
            for material in production.material_line_ids.sorted(
                lambda material: material.id
            )
        ]

        setup_line.action_select_fabric_editor()
        setup.write({
            'editor_line_ids': [(0, 0, {
                'product_id': fixture['manual_fabric'].id,
                'qty': 7.25,
                'product_uom_id': fixture['manual_fabric'].uom_id.id,
            })],
        })
        action = setup.action_cancel_active_editor()

        self.assertEqual(action['type'], 'ir.actions.client')
        self.assertEqual(action['tag'], 'furniture_tailoring_setup_refresh')
        self.assertFalse(setup.active_editor_mode)
        self.assertFalse(setup.active_production_line_id)
        self.assertFalse(setup.editor_line_ids)
        production.invalidate_recordset(['tailoring_material_revision'])
        production_line.invalidate_recordset([
            'tailoring_material_allocation_ids',
            'tailoring_fabric_configured',
        ])
        self.assertEqual(
            production.tailoring_material_revision,
            baseline_revision,
        )
        self.assertEqual(
            production_line.tailoring_material_allocation_ids.ids,
            baseline_allocation_ids,
        )
        self.assertFalse(production_line.tailoring_fabric_configured)
        self.assertEqual([
            (
                material.id,
                material.product_id.id,
                material.stage,
                material.qty_needed,
            )
            for material in production.material_line_ids.sorted(
                lambda material: material.id
            )
        ], baseline_materials)

    def test_inline_apply_changes_only_selected_piece_and_kind(self):
        fixture = self._create_fixture()
        production = fixture['production']
        target_line = fixture['lines'][0]
        other_line = fixture['lines'][1]
        self._apply_editor(target_line, 'takawe', [
            (fixture['manual_takawe'], 1.5, False),
        ])
        self._apply_editor(other_line, 'fabric', [
            (fixture['second_fabric'], 2.25, False),
        ])
        target_takawe = [
            (allocation.product_id.id, allocation.qty)
            for allocation in target_line.tailoring_material_allocation_ids.filtered(
                lambda allocation: allocation.material_kind == 'takawe'
            )
        ]
        other_fabric = [
            (allocation.product_id.id, allocation.qty)
            for allocation in other_line.tailoring_material_allocation_ids.filtered(
                lambda allocation: allocation.material_kind == 'fabric'
            )
        ]
        other_material_line_ids = production.material_line_ids.filtered(
            lambda material: material.production_line_id == other_line
        ).ids

        setup = self._open_setup(production)
        setup_line = setup.line_ids.filtered(
            lambda row: row.production_line_id == target_line
        ).ensure_one()
        setup_count = self.env[
            'furniture.mrp.tailoring.setup.wizard'
        ].search_count([])
        setup_line.action_select_fabric_editor()
        setup.write({
            'editor_line_ids': [(5, 0, 0), (0, 0, {
                'product_id': fixture['manual_fabric'].id,
                'qty': 4.75,
                'product_uom_id': fixture['manual_fabric'].uom_id.id,
            })],
        })
        action = setup.action_apply_active_editor()

        self.assertEqual(action['type'], 'ir.actions.client')
        self.assertEqual(action['tag'], 'furniture_tailoring_setup_refresh')
        self.assertEqual(
            self.env['furniture.mrp.tailoring.setup.wizard'].search_count([]),
            setup_count,
            'Inline save must not create or stack another setup dialog.',
        )
        target_line.invalidate_recordset([
            'tailoring_material_allocation_ids',
            'tailoring_fabric_configured',
            'tailoring_takawe_configured',
        ])
        other_line.invalidate_recordset([
            'tailoring_material_allocation_ids',
            'tailoring_fabric_configured',
        ])
        target_fabric = target_line.tailoring_material_allocation_ids.filtered(
            lambda allocation: allocation.material_kind == 'fabric'
        )
        self.assertEqual(target_fabric.product_id, fixture['manual_fabric'])
        self.assertAlmostEqual(sum(target_fabric.mapped('qty')), 4.0)
        self.assertEqual([
            (allocation.product_id.id, allocation.qty)
            for allocation in target_line.tailoring_material_allocation_ids.filtered(
                lambda allocation: allocation.material_kind == 'takawe'
            )
        ], target_takawe)
        self.assertEqual([
            (allocation.product_id.id, allocation.qty)
            for allocation in other_line.tailoring_material_allocation_ids.filtered(
                lambda allocation: allocation.material_kind == 'fabric'
            )
        ], other_fabric)
        self.assertEqual(
            production.material_line_ids.filtered(
                lambda material: material.production_line_id == other_line
            ).ids,
            other_material_line_ids,
            'Saving one material kind must not rebuild unrelated recipe lines.',
        )
        self.assertTrue(target_line.tailoring_fabric_configured)
        self.assertTrue(target_line.tailoring_takawe_configured)
        self.assertTrue(other_line.tailoring_fabric_configured)
        self.assertFalse(setup.active_editor_mode)
        self.assertFalse(setup.editor_line_ids)

    def test_fabric_selection_replaces_only_auto_fabric_and_survives_refresh(self):
        fixture = self._create_fixture()
        line = fixture['lines'][0]
        original = self._tailoring_materials(line)
        original_other_qty = sum(
            original.filtered(
                lambda material: material.product_id == fixture['other_material']
            ).mapped('qty_needed')
        )
        original_takawe_qty = sum(
            self._upholstery_materials(line).filtered(
                lambda material: material.product_id == fixture['auto_takawe']
            ).mapped('qty_needed')
        )

        self._apply_editor(line, 'fabric', [
            (fixture['manual_fabric'], 99.0, False),
        ])

        self.assertTrue(line.tailoring_fabric_configured)
        self.assertFalse(line.tailoring_takawe_configured)
        allocation_relation = self._allocation_relation_field()
        kind_field = self._allocation_kind_field()
        self.assertEqual(
            len(line[allocation_relation].filtered(
                lambda allocation: allocation[kind_field] == 'fabric'
            )),
            1,
        )
        for refresh in (False, True):
            if refresh:
                fixture['production']._refresh_material_lines_for_stage_plan()
            materials = self._tailoring_materials(line)
            self.assertNotIn(fixture['auto_fabric'], materials.mapped('product_id'))
            self.assertEqual(
                set(materials.filtered(
                    lambda material: material.product_id
                    == fixture['manual_fabric']
                ).mapped('product_id').ids),
                {fixture['manual_fabric'].id},
            )
            self.assertAlmostEqual(
                sum(materials.filtered(
                    lambda material: material.product_id == fixture['manual_fabric']
                ).mapped('qty_needed')),
                4.0,
            )
            self.assertAlmostEqual(
                sum(materials.filtered(
                    lambda material: material.product_id == fixture['other_material']
                ).mapped('qty_needed')),
                original_other_qty,
            )
            self.assertAlmostEqual(
                sum(self._upholstery_materials(line).filtered(
                    lambda material: material.product_id == fixture['auto_takawe']
                ).mapped('qty_needed')),
                original_takawe_qty,
            )
        source_recipe_products = fixture[
            'boms'
        ][0].furniture_stage_material_line_ids.mapped('product_id')
        self.assertIn(fixture['auto_fabric'], source_recipe_products)
        self.assertNotIn(fixture['manual_fabric'], source_recipe_products)
        self.assertNotIn(fixture['second_fabric'], source_recipe_products)

    def test_takawe_selection_replaces_only_auto_takawe(self):
        fixture = self._create_fixture()
        line = fixture['lines'][1]
        tailoring_materials = self._tailoring_materials(line)
        upholstery_materials = self._upholstery_materials(line)
        self.assertIn(
            fixture['auto_takawe'],
            upholstery_materials.mapped('product_id'),
        )
        self.assertNotIn(
            fixture['auto_takawe'],
            tailoring_materials.mapped('product_id'),
        )
        source_takawe = fixture['boms'][1].furniture_stage_material_line_ids.filtered(
            lambda material: material.product_id == fixture['auto_takawe']
        ).ensure_one()
        self.assertEqual(source_takawe.stage, 'upholstery')
        original_fabric_qty = sum(
            tailoring_materials.filtered(
                lambda material: material.product_id == fixture['auto_fabric']
            ).mapped('qty_needed')
        )
        original_other_qty = sum(
            tailoring_materials.filtered(
                lambda material: material.product_id == fixture['other_material']
            ).mapped('qty_needed')
        )

        self._apply_editor(line, 'takawe', [
            (fixture['manual_takawe'], 4.0, False),
        ])
        fixture['production']._refresh_material_lines_for_stage_plan()

        self.assertFalse(line.tailoring_fabric_configured)
        self.assertTrue(line.tailoring_takawe_configured)
        upholstery_materials = self._upholstery_materials(line)
        self.assertNotIn(
            fixture['auto_takawe'],
            upholstery_materials.mapped('product_id'),
        )
        self.assertAlmostEqual(
            sum(upholstery_materials.filtered(
                lambda material: material.product_id == fixture['manual_takawe']
            ).mapped('qty_needed')),
            4.0,
        )
        tailoring_materials = self._tailoring_materials(line)
        self.assertNotIn(
            fixture['manual_takawe'],
            tailoring_materials.mapped('product_id'),
        )
        self.assertAlmostEqual(
            sum(tailoring_materials.filtered(
                lambda material: material.product_id == fixture['auto_fabric']
            ).mapped('qty_needed')),
            original_fabric_qty,
        )
        self.assertAlmostEqual(
            sum(tailoring_materials.filtered(
                lambda material: material.product_id == fixture['other_material']
            ).mapped('qty_needed')),
            original_other_qty,
        )

    def test_imported_placeholders_use_authoritative_replacement_stages(self):
        script_path = (
            Path(__file__).resolve().parents[1]
            / 'scripts'
            / 'import_yasser3_workbook_recipes.py'
        )
        spec = importlib.util.spec_from_file_location(
            'test_import_yasser3_workbook_recipes',
            script_path,
        )
        importer = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(importer)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Regression Model'
        sheet.cell(4, 4).value = 'كنبة كبيرة'
        sheet.cell(5, 2).value = 'قماش'
        for row, material, qty in (
            (6, 'قماش', 2.0),
            (7, 'مشجر', 1.0),
        ):
            sheet.cell(row, 4).value = material
            sheet.cell(row, 5).value = 'متر'
            sheet.cell(row, 6).value = qty

        with patch.object(importer, 'load_workbook', return_value=workbook):
            models, _issues = importer.parse_workbook('unused.xlsx')
        imported_lines = models[0]['recipes'][0]['lines']
        self.assertEqual(
            [
                (
                    importer.TAILORING_MATERIAL_KINDS[
                        importer.key(item['material'])
                    ],
                    item['stage'],
                )
                for item in imported_lines
            ],
            [('fabric', 'tailoring'), ('takawe', 'upholstery')],
        )

        fixture = self._create_fixture(line_quantities=(1.0, 1.0, 1.0))
        line = fixture['lines'][0]
        fixture['auto_fabric'].write({'name': 'قماش'})
        fixture['auto_takawe'].write({'name': 'مشجر'})
        generic_products = fixture['auto_fabric'] | fixture['auto_takawe']

        original_placeholders = (
            fixture['production'].material_line_ids.filtered(
                lambda material: (
                    material.production_line_id == line
                    and material.product_id in generic_products
                )
            )
        )
        self.assertEqual(
            {
                (material.product_id.id, material.stage)
                for material in original_placeholders
            },
            {
                (fixture['auto_fabric'].id, 'tailoring'),
                (fixture['auto_takawe'].id, 'upholstery'),
            },
        )

        self._apply_editor(line, 'fabric', [
            (fixture['manual_fabric'], 5.5, False),
        ])
        self._apply_editor(line, 'takawe', [
            (fixture['manual_takawe'], 3.25, False),
        ])
        fixture['production']._refresh_material_lines_for_stage_plan()

        materials = fixture['production'].material_line_ids.filtered(
            lambda material: material.production_line_id == line
        )
        self.assertFalse(
            materials.filtered(
                lambda material: material.product_id in generic_products
            ),
            'Selected materials must replace, not duplicate, generic '
            'placeholders.',
        )
        selected_materials = materials.filtered('tailoring_allocation_id')
        self.assertEqual(
            {
                (
                    material.tailoring_allocation_id.material_kind,
                    material.product_id.id,
                    material.stage,
                    material.qty_needed,
                )
                for material in selected_materials
            },
            {
                ('fabric', fixture['manual_fabric'].id, 'tailoring', 2.0),
                ('takawe', fixture['manual_takawe'].id, 'upholstery', 3.25),
            },
        )

    def test_importer_splits_side_by_side_outputs_into_independent_recipes(self):
        script_path = (
            Path(__file__).resolve().parents[1]
            / 'scripts'
            / 'import_yasser3_workbook_recipes.py'
        )
        spec = importlib.util.spec_from_file_location(
            'test_import_yasser3_split_outputs',
            script_path,
        )
        importer = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(importer)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Regression Model'
        sheet.cell(4, 4).value = 'كنبه كبيره'
        sheet.cell(4, 5).value = 'شيزلونج'
        sheet.cell(4, 10).value = 'فوتيه'
        sheet.cell(5, 2).value = 'أولاً: تقديم'
        sheet.cell(6, 4).value = 'خشب زان 5'
        sheet.cell(6, 5).value = 'متر مكعب'
        sheet.cell(6, 6).value = 0.06
        sheet.cell(6, 7).value = 23250
        sheet.cell(6, 10).value = 'CO16m'
        sheet.cell(6, 11).value = 'لوح'
        sheet.cell(6, 12).value = 0.13
        sheet.cell(6, 13).value = 1230

        with patch.object(importer, 'load_workbook', return_value=workbook):
            models, issues = importer.parse_workbook('unused.xlsx')

        recipes = models[0]['recipes']
        self.assertEqual(
            [(recipe['product'], recipe['kit_qty']) for recipe in recipes],
            [
                ('كنبة كبيرة', 1.0),
                ('شازلونج', 1.0),
                ('فوتيه', 2.0),
            ],
        )
        self.assertNotIn(
            'كنبة كبيرة + شازلونج',
            [recipe['product'] for recipe in recipes],
        )
        self.assertEqual(recipes[0]['lines'], recipes[1]['lines'])
        self.assertIsNot(recipes[0]['lines'], recipes[1]['lines'])
        self.assertIsNot(recipes[0]['lines'][0], recipes[1]['lines'][0])
        self.assertEqual(recipes[0]['lines'][0]['source_qty'], 0.06)
        self.assertEqual(recipes[0]['lines'][0]['qty'], 0.06)
        self.assertEqual(recipes[1]['lines'][0]['qty'], 0.06)
        self.assertEqual(recipes[2]['lines'][0]['source_qty'], 0.13)
        self.assertEqual(recipes[2]['lines'][0]['qty'], 0.13)
        self.assertEqual(
            recipes[0]['lines'][0]['qty'],
            0.06,
        )
        self.assertEqual(
            recipes[1]['lines'][0]['qty'],
            0.06,
        )
        self.assertEqual(
            recipes[2]['lines'][0]['qty'],
            0.13,
        )
        self.assertEqual(importer.kit_material_totals(models), {
            ('regression model', 'priming', 'خشب زان 5 سم', 'cubic_meter'): 0.12,
            ('regression model', 'priming', 'co 16m', 'board'): 0.26,
        })
        self.assertEqual(
            importer.build_summary(models, issues),
            {
                'models': 1,
                'recipes': 3,
                'stage_lines': 3,
                'kit_lines': 3,
                'missing_qty': 0,
                'defaulted_uom': 0,
                'inferred_uom': 0,
                'products': ['شازلونج', 'فوتيه', 'كنبة كبيرة'],
                'materials': 2,
            },
        )

    def test_allocation_rejects_zero_or_negative_quantity(self):
        fixture = self._create_fixture()
        for qty in (0.0, -1.0):
            with self.assertRaises(ValidationError):
                self._create_allocation(
                    fixture['lines'][0],
                    'fabric',
                    fixture['manual_fabric'],
                    qty,
                )

    def test_allocation_forces_meter_when_another_uom_is_supplied(self):
        fixture = self._create_fixture()
        weight_uom = self.env.ref('uom.product_uom_kgm')
        meter = self.env.ref('furniture_mrp.furniture_uom_meter')

        allocation = self._create_allocation(
            fixture['lines'][0],
            'fabric',
            fixture['manual_fabric'],
            4.0,
            uom=weight_uom,
        )

        self.assertEqual(allocation.product_uom_id, meter)

    def test_direct_fabric_allocation_cannot_change_recipe_meters(self):
        fixture = self._create_fixture()
        allocation = self._create_allocation(
            fixture['lines'][0],
            'fabric',
            fixture['manual_fabric'],
            4.0,
        )

        with self.assertRaisesRegex(ValidationError, 'ثابت من الريسيبي'):
            allocation.write({'qty': 3.0})
        with self.assertRaisesRegex(ValidationError, 'نافذة الأقمشة'):
            allocation.unlink()

    def test_allocation_and_material_line_always_use_meter(self):
        fixture = self._create_fixture()
        line = fixture['lines'][0]
        units = self.env.ref('uom.product_uom_unit')
        dozen = self.env.ref('uom.product_uom_dozen')
        meter = self.env.ref('furniture_mrp.furniture_uom_meter')
        fixture['manual_fabric'].write({
            'uom_id': units.id,
            'uom_po_id': units.id,
        })
        self._apply_editor(line, 'fabric', [
            (fixture['manual_fabric'], 1.0, dozen),
        ])
        allocation = line.tailoring_material_allocation_ids.filtered(
            lambda item: item.material_kind == 'fabric'
        ).ensure_one()

        self.assertAlmostEqual(allocation.qty, 4.0)
        self.assertEqual(allocation.product_uom_id, meter)
        material_line = fixture['production'].material_line_ids.filtered(
            lambda material: material.tailoring_allocation_id == allocation
        ).ensure_one()
        self.assertAlmostEqual(material_line.qty_needed, 4.0)
        self.assertEqual(material_line.product_uom_id, meter)
        setup_row = self._open_setup(fixture['production']).line_ids.filtered(
            lambda row: row.production_line_id == line
        ).ensure_one()
        self.assertEqual(
            setup_row.fabric_summary,
            '%s%s4 %s' % (
                fixture['manual_fabric'].display_name,
                '\u00a0\u00a0•\u00a0\u00a0',
                meter.display_name,
            ),
        )

    def test_editor_merges_same_product_into_single_meter_allocation(self):
        fixture = self._create_fixture()
        line = fixture['lines'][0]
        units = self.env.ref('uom.product_uom_unit')
        dozen = self.env.ref('uom.product_uom_dozen')
        meter = self.env.ref('furniture_mrp.furniture_uom_meter')
        fixture['manual_takawe'].write({
            'uom_id': units.id,
            'uom_po_id': units.id,
        })

        self._apply_editor(line, 'takawe', [
            (fixture['manual_takawe'], 1.0, dozen),
            (fixture['manual_takawe'], 2.0, units),
        ])

        allocations = line.tailoring_material_allocation_ids.filtered(
            lambda allocation: allocation.material_kind == 'takawe'
        )
        self.assertEqual(len(allocations), 1)
        self.assertEqual(
            {
                (allocation.product_uom_id.id, allocation.qty)
                for allocation in allocations
            },
            {(meter.id, 3.0)},
        )
        material_lines = fixture['production'].material_line_ids.filtered(
            lambda material: material.tailoring_allocation_id in allocations
        )
        self.assertEqual(
            {
                (material.product_uom_id.id, material.qty_needed)
                for material in material_lines
            },
            {(meter.id, 3.0)},
        )
        setup_row = self._open_setup(fixture['production']).line_ids.filtered(
            lambda row: row.production_line_id == line
        ).ensure_one()
        quantity_separator = '\u00a0\u00a0•\u00a0\u00a0'
        self.assertEqual(
            setup_row.takawe_summary,
            '%s%s3 %s · مقاس 45' % (
                fixture['manual_takawe'].display_name,
                quantity_separator,
                meter.display_name,
            ),
        )

    def test_both_classified_material_types_work_in_both_editors(self):
        fixture = self._create_fixture()
        line = fixture['lines'][0]

        self._apply_editor(line, 'fabric', [
            (fixture['manual_takawe'], 99.0, False),
        ])
        fabric_allocation = line.tailoring_material_allocation_ids.filtered(
            lambda allocation: allocation.material_kind == 'fabric'
        ).ensure_one()
        self.assertEqual(fabric_allocation.product_id, fixture['manual_takawe'])
        self.assertAlmostEqual(fabric_allocation.qty, 4.0)

        self._apply_editor(line, 'takawe', [
            (fixture['manual_fabric'], 2.75, False),
        ])
        takawe_allocation = line.tailoring_material_allocation_ids.filtered(
            lambda allocation: allocation.material_kind == 'takawe'
        ).ensure_one()
        self.assertEqual(takawe_allocation.product_id, fixture['manual_fabric'])
        self.assertAlmostEqual(takawe_allocation.qty, 2.75)

    def test_allocation_rejects_finished_or_unclassified_product(self):
        fixture = self._create_fixture()

        with self.assertRaises(ValidationError):
            self._create_allocation(
                fixture['lines'][0],
                'fabric',
                fixture['products'][1],
                1.0,
            )
        with self.assertRaisesRegex(ValidationError, 'صنّف الخامة'):
            self._create_allocation(
                fixture['lines'][0],
                'fabric',
                fixture['other_material'],
                1.0,
            )

    def test_allocation_rejects_production_line_from_another_order(self):
        fixture = self._create_fixture()
        foreign = self._create_fixture(line_quantities=(1.0, 1.0, 1.0))
        Allocation = self.env[
            'furniture.mrp.tailoring.material.allocation'
        ]
        if 'production_id' not in Allocation._fields:
            self.fail(
                'The allocation needs an order identity so a forged line from '
                'another production can be rejected.'
            )

        with self.assertRaises(ValidationError):
            self._create_allocation(
                foreign['lines'][0],
                'fabric',
                fixture['manual_fabric'],
                1.0,
                production=fixture['production'],
            )

    def test_setup_access_and_closed_state_are_locked(self):
        fixture = self._create_fixture()
        production = fixture['production']

        self.assertTrue(
            production.with_user(
                self.manager_user
            ).action_open_tailoring_material_setup()
        )
        with self.assertRaises(AccessError):
            production.with_user(
                self.regular_user
            ).action_open_tailoring_material_setup()

        for closed_state in ('done', 'cancelled'):
            closed_production = (
                production
                if closed_state == 'done'
                else self._create_fixture()['production']
            )
            closed_production.with_context(
                furniture_skip_material_refresh=True,
                furniture_skip_stage_plan_sync=True,
            ).write({'state': closed_state})
            with self.assertRaises(UserError):
                closed_production.action_open_tailoring_material_setup()

    def test_active_store_request_locks_open_editor_and_apply(self):
        for request_state in ('pending', 'approved', 'started'):
            fixture = self._create_fixture()
            production = fixture['production']
            setup = self._open_setup(production)
            setup_line = setup.line_ids.filtered(
                lambda row: row.production_line_id == fixture['lines'][0]
            ).ensure_one()
            editor_action = setup_line.action_open_fabric_editor()
            editor_params = editor_action['params']
            editor = self.env[editor_params['res_model']].browse(
                editor_params['res_id']
            )
            request = self._create_store_request(fixture, request_state)
            self.assertTrue(request)

            with self.assertRaises(UserError):
                production.action_open_tailoring_material_setup()
            with self.assertRaises(UserError):
                setup_line.action_open_fabric_editor()
            with self.assertRaises(UserError):
                editor.action_apply()

    def test_rejected_or_cancelled_store_request_does_not_lock_setup(self):
        for request_state in ('rejected', 'cancelled'):
            fixture = self._create_fixture()
            request = self._create_store_request(fixture, request_state)

            self.assertTrue(request)
            self.assertTrue(
                fixture['production'].action_open_tailoring_material_setup()
            )

    def test_direct_allocation_mutation_obeys_setup_lock(self):
        fixture = self._create_fixture()
        allocation = self._create_allocation(
            fixture['lines'][0],
            'fabric',
            fixture['manual_fabric'],
            4.0,
        )
        self._create_store_request(fixture, 'pending')

        with self.assertRaises(UserError):
            allocation.write({'qty': 3.0})
        with self.assertRaises(UserError):
            allocation.unlink()
        with self.assertRaises(UserError):
            self._create_allocation(
                fixture['lines'][1],
                'fabric',
                fixture['second_fabric'],
                1.0,
            )

    def test_supervisor_cannot_open_or_mutate_manager_setup(self):
        fixture = self._create_fixture()
        line = fixture['lines'][0]
        production = fixture['production']

        with self.assertRaises(AccessError):
            production.with_user(
                self.supervisor_user
            ).action_open_tailoring_material_setup()

        # A forged editable transient must not turn the read-only supervisor
        # role into a persistent material/image write capability.
        setup = self.env[
            'furniture.mrp.tailoring.setup.wizard'
        ].with_user(self.supervisor_user).create({
            'production_id': production.id,
            'line_ids': [(0, 0, {
                'sequence': line.sequence,
                'production_line_id': line.id,
            })],
        })
        setup_line = setup.line_ids.ensure_one()

        with self.assertRaises(AccessError):
            setup_line.action_select_fabric_editor()
        with self.assertRaises(AccessError):
            setup_line.upload_piece_image(TEST_IMAGE_2)
        with self.assertRaises(AccessError):
            setup.upload_order_image(TEST_IMAGE_2)
        with self.assertRaises(AccessError):
            setup_line.save_piece_note('غير مسموح')

        line.invalidate_recordset(['batch_image_1920', 'batch_image_token'])
        production.invalidate_recordset([
            'tailoring_set_image_1920',
            'tailoring_set_image_token',
        ])
        self.assertFalse(line.tailoring_material_allocation_ids)
        self.assertFalse(line.batch_image_1920)
        self.assertFalse(line.batch_image_token)
        self.assertFalse(line.kit_piece_note)
        self.assertFalse(production.tailoring_set_image_1920)
        self.assertFalse(production.tailoring_set_image_token)

    def test_editor_rejects_line_changed_after_open(self):
        fixture = self._create_fixture(line_quantities=(4.0, 1.0, 1.0))
        line = fixture['lines'][0]
        editor = self._open_editor(line, 'fabric')
        editor.write({
            'piece_size': '50',
            'line_ids': [(0, 0, {
                'product_id': fixture['manual_fabric'].id,
                'qty': 2.0,
                'product_uom_id': fixture['manual_fabric'].uom_id.id,
            })],
        })
        line._split_for_partial_quantity(1.5)

        with self.assertRaises(UserError):
            editor.action_apply()

    def test_partial_split_scales_and_copies_material_allocations(self):
        fixture = self._create_fixture(line_quantities=(4.0, 1.0, 1.0))
        line = fixture['lines'][0]
        self._apply_editor(line, 'fabric', [
            (fixture['manual_fabric'], 10.0, False),
        ])

        remaining = line._split_for_partial_quantity(1.5)
        self.assertTrue(remaining)
        self.assertAlmostEqual(line.product_qty, 1.5)
        self.assertAlmostEqual(remaining.product_qty, 2.5)
        self.assertTrue(line.tailoring_fabric_configured)
        self.assertTrue(remaining.tailoring_fabric_configured)

        Allocation = self.env[
            'furniture.mrp.tailoring.material.allocation'
        ]
        kind_field = self._allocation_kind_field()
        kept_allocations = Allocation.search([
            ('production_line_id', '=', line.id),
            (kind_field, '=', 'fabric'),
        ])
        remaining_allocations = Allocation.search([
            ('production_line_id', '=', remaining.id),
            (kind_field, '=', 'fabric'),
        ])
        self.assertEqual(len(kept_allocations), 1)
        self.assertEqual(len(remaining_allocations), 1)
        self.assertAlmostEqual(sum(kept_allocations.mapped('qty')), 3.0)
        self.assertAlmostEqual(sum(remaining_allocations.mapped('qty')), 5.0)
        self.assertEqual(
            kept_allocations.product_id,
            fixture['manual_fabric'],
        )
        self.assertEqual(
            remaining_allocations.product_id,
            fixture['manual_fabric'],
        )

    def test_configured_split_lines_are_not_consolidated(self):
        fixture = self._create_fixture(line_quantities=(1.0, 1.0, 1.0))
        production = fixture['production']
        source = fixture['lines'][0]
        duplicate = source.with_context(
            furniture_skip_line_consolidation=True,
            furniture_skip_material_refresh=True,
            furniture_skip_stage_plan_sync=True,
        ).copy({
            'sequence': source.sequence + 1,
            'product_qty': source.product_qty,
        })
        for line in source | duplicate:
            self._apply_editor(line, 'fabric', [
                (fixture['manual_fabric'], 2.0, False),
            ])

        production._consolidate_equivalent_production_lines()
        (source | duplicate).invalidate_recordset([
            'active',
            'consolidated_into_line_id',
        ])

        self.assertTrue(source.active)
        self.assertTrue(duplicate.active)
        self.assertFalse(source.consolidated_into_line_id)
        self.assertFalse(duplicate.consolidated_into_line_id)

    def test_order_image_is_shared_without_writing_piece_or_product_images(self):
        fixture = self._create_fixture()
        production = fixture['production']
        product_images = {
            line.product_id.id: line.product_id.image_1920
            for line in fixture['lines']
        }
        setup = self._open_setup(production)
        stale_setup = self._open_setup(production)

        self.assertFalse(setup.order_image_128)
        self.assertFalse(setup.has_order_image)
        result = setup.with_user(self.manager_user).upload_order_image(
            TEST_IMAGE_2
        )
        self.assertTrue(result['image_token'])

        production.invalidate_recordset([
            'tailoring_set_image_1920',
            'tailoring_set_image_token',
        ])
        setup.invalidate_recordset(['order_image_128', 'has_order_image'])
        fixture['lines'].invalidate_recordset([
            'batch_image_1920',
            'batch_image_token',
        ])
        self.assertEqual(production.tailoring_set_image_1920, TEST_IMAGE_2)
        self.assertEqual(
            production.tailoring_set_image_token,
            result['image_token'],
        )
        self.assertTrue(setup.has_order_image)
        self.assertEqual(setup.order_image_128, TEST_IMAGE_2)
        # Match the web client's field order and bin-size context.  The
        # presence Boolean must never cause a human-readable size such as
        # ``5.15 Kb`` to be assigned back into fields.Image as base64.
        web_values = setup.with_context(bin_size=True).read([
            'has_order_image',
            'order_image_128',
        ])[0]
        self.assertTrue(web_values['has_order_image'])
        self.assertTrue(web_values['order_image_128'])
        self.assertFalse(any(fixture['lines'].mapped('batch_image_1920')))
        self.assertFalse(any(fixture['lines'].mapped('batch_image_token')))
        self.assertEqual(len(self._order_image_attachments(production)), 1)
        for line in fixture['lines']:
            line.product_id.invalidate_recordset(['image_1920'])
            self.assertEqual(
                line.product_id.image_1920,
                product_images[line.product_id.id],
            )

        with self.assertRaises(UserError):
            stale_setup.with_user(self.manager_user).upload_order_image(
                TEST_IMAGE
            )

        viewer = self.env[
            'furniture.mrp.tailoring.setup.wizard'
        ].with_user(self.supervisor_user).create({
            'production_id': production.id,
            'view_only': True,
            'viewer_stage': 'tailoring',
            'line_ids': [(0, 0, {
                'sequence': line.sequence,
                'production_line_id': line.id,
            }) for line in fixture['lines']],
        })
        self.assertTrue(viewer.has_order_image)
        self.assertEqual(viewer.order_image_128, TEST_IMAGE_2)
        with self.assertRaises(AccessError):
            viewer.upload_order_image(TEST_IMAGE)

    def test_order_image_gallery_adds_previews_removes_and_checks_access(self):
        fixture = self._create_fixture()
        production = fixture['production']
        setup = self._open_setup(production)
        stale_setup = self._open_setup(production)

        first = setup.with_user(self.manager_user).add_order_image(
            TEST_IMAGE,
            'main.png',
        )
        self.assertEqual([image['key'] for image in first['images']], ['main'])
        second = setup.with_user(self.manager_user).add_order_image(
            TEST_IMAGE_2,
            'extra.png',
        )
        self.assertEqual(len(second['images']), 2)
        self.assertEqual(second['images'][0]['key'], 'main')
        self.assertTrue(second['images'][1]['key'].startswith('attachment:'))

        production.invalidate_recordset([
            'tailoring_set_image_1920',
            'tailoring_set_image_token',
            'tailoring_set_image_attachment_ids',
        ])
        self.assertEqual(production.tailoring_set_image_1920, TEST_IMAGE)
        extra_attachment = (
            production.sudo().tailoring_set_image_attachment_ids.ensure_one()
        )
        self.assertEqual(extra_attachment.datas, TEST_IMAGE_2)
        self.assertEqual(extra_attachment.name, 'extra.png')
        self.assertEqual(extra_attachment.res_model, production._name)
        self.assertEqual(extra_attachment.res_id, production.id)

        with self.assertRaises(UserError):
            stale_setup.with_user(self.manager_user).add_order_image(
                TEST_IMAGE_2,
                'stale.png',
            )

        viewer = self.env[
            'furniture.mrp.tailoring.setup.wizard'
        ].with_user(self.supervisor_user).create({
            'production_id': production.id,
            'view_only': True,
            'viewer_stage': 'tailoring',
            'line_ids': [(0, 0, {
                'sequence': line.sequence,
                'production_line_id': line.id,
            }) for line in fixture['lines']],
        })
        self.assertEqual(len(viewer.get_order_image_gallery()), 2)
        extra_attachment.with_user(self.supervisor_user).check_access('read')
        with self.assertRaises(AccessError):
            viewer.add_order_image(TEST_IMAGE, 'forbidden.png')
        with self.assertRaises(AccessError):
            viewer.remove_order_image('main')

        extra_key = second['images'][1]['key']
        after_extra_delete = setup.with_user(
            self.manager_user
        ).remove_order_image(extra_key)
        self.assertEqual(
            [image['key'] for image in after_extra_delete['images']],
            ['main'],
        )
        self.assertFalse(production.sudo().tailoring_set_image_attachment_ids)

        setup.with_user(self.manager_user).add_order_image(
            TEST_IMAGE_2,
            'promote.png',
        )
        promoted = setup.with_user(self.manager_user).remove_order_image('main')
        production.invalidate_recordset([
            'tailoring_set_image_1920',
            'tailoring_set_image_token',
            'tailoring_set_image_attachment_ids',
        ])
        self.assertEqual([image['key'] for image in promoted['images']], ['main'])
        self.assertEqual(production.tailoring_set_image_1920, TEST_IMAGE_2)
        self.assertFalse(production.sudo().tailoring_set_image_attachment_ids)

        with patch.object(
            mrp_tailoring_material_setup,
            'TAILORING_ORDER_IMAGE_LIMIT',
            1,
        ), self.assertRaises(ValidationError):
            setup.with_user(self.manager_user).add_order_image(
                TEST_IMAGE,
                'too-many.png',
            )

    def test_piece_note_saves_on_exact_row_and_is_readonly_for_supervisor(self):
        fixture = self._create_fixture()
        production = fixture['production']
        line = fixture['lines'][0]
        other_line = fixture['lines'][1]
        setup = self._open_setup(production)
        stale_setup = self._open_setup(production)
        setup_row = setup.line_ids.filtered(
            lambda row: row.production_line_id == line
        ).ensure_one()

        result = setup_row.with_user(self.manager_user).save_piece_note(
            'ارتفاع الظهر 45 سم'
        )
        self.assertEqual(result['note'], 'ارتفاع الظهر 45 سم')
        line.invalidate_recordset(['kit_piece_note'])
        other_line.invalidate_recordset(['kit_piece_note'])
        self.assertEqual(line.kit_piece_note, 'ارتفاع الظهر 45 سم')
        self.assertFalse(other_line.kit_piece_note)

        stale_row = stale_setup.line_ids.filtered(
            lambda row: row.production_line_id == line
        ).ensure_one()
        with self.assertRaises(UserError):
            stale_row.with_user(self.manager_user).save_piece_note(
                'ملاحظة قديمة'
            )

        viewer = self.env[
            'furniture.mrp.tailoring.setup.wizard'
        ].with_user(self.supervisor_user).create({
            'production_id': production.id,
            'view_only': True,
            'viewer_stage': 'tailoring',
            'line_ids': [(0, 0, {
                'sequence': line.sequence,
                'production_line_id': line.id,
            })],
        })
        viewer_row = viewer.line_ids.ensure_one()
        self.assertEqual(viewer_row.piece_note, 'ارتفاع الظهر 45 سم')
        with self.assertRaises(AccessError):
            viewer_row.save_piece_note('تعديل غير مسموح')

    def test_direct_piece_image_upload_targets_exact_line_and_keeps_material_draft(self):
        fixture = self._create_fixture()
        line = fixture['lines'][0]
        other_line = fixture['lines'][1]
        product = line.product_id
        self.assertEqual(product.image_1920, TEST_IMAGE)
        self.assertFalse(line.batch_image_1920)

        setup = self._open_setup(
            fixture['production'],
            user=self.manager_user,
        )
        setup_line = setup.line_ids.filtered(
            lambda row: row.production_line_id == line
        ).ensure_one()
        self.assertFalse(setup_line.image_128)
        self.assertFalse(setup_line.has_image)

        setup_line.action_select_fabric_editor()
        setup.write({
            'editor_line_ids': [(5, 0, 0), (0, 0, {
                'sequence': 10,
                'product_id': fixture['manual_fabric'].id,
                'qty': 3.75,
                'product_uom_id': fixture['manual_fabric'].uom_id.id,
            })],
        })
        draft_signature = [
            (
                draft.product_id.id,
                draft.qty,
                draft.product_uom_id.id,
            )
            for draft in setup.editor_line_ids
        ]

        result = setup_line.with_user(
            self.manager_user
        ).upload_piece_image(TEST_IMAGE_2)
        self.assertTrue(result)

        product.invalidate_recordset(['image_1920'])
        line.invalidate_recordset(['batch_image_1920', 'batch_image_token'])
        other_line.invalidate_recordset([
            'batch_image_1920',
            'batch_image_token',
        ])
        setup.invalidate_recordset([
            'active_editor_mode',
            'active_production_line_id',
            'editor_line_ids',
        ])
        self.assertEqual(product.image_1920, TEST_IMAGE)
        self.assertEqual(line.batch_image_1920, TEST_IMAGE_2)
        self.assertTrue(line.batch_image_token)
        self.assertFalse(other_line.batch_image_1920)
        self.assertFalse(other_line.batch_image_token)
        self.assertEqual(len(self._piece_image_attachments(line)), 1)
        self.assertEqual(setup.active_editor_mode, 'fabric')
        self.assertEqual(setup.active_production_line_id, line)
        self.assertEqual([
            (
                draft.product_id.id,
                draft.qty,
                draft.product_uom_id.id,
            )
            for draft in setup.editor_line_ids
        ], draft_signature)

    def test_direct_piece_image_replacement_accepts_data_url_and_rotates_token(self):
        fixture = self._create_fixture()
        line = fixture['lines'][0]
        product = line.product_id
        setup_line = self._open_setup(
            fixture['production']
        ).line_ids.filtered(
            lambda row: row.production_line_id == line
        ).ensure_one()

        self.assertTrue(setup_line.upload_piece_image(TEST_IMAGE))
        line.invalidate_recordset(['batch_image_1920', 'batch_image_token'])
        first_token = line.batch_image_token
        first_attachment = self._piece_image_attachments(line).ensure_one()
        first_checksum = first_attachment.checksum

        data_url = 'data:image/png;base64,%s' % TEST_IMAGE_2.decode()
        self.assertTrue(setup_line.upload_piece_image(data_url))

        line.invalidate_recordset(['batch_image_1920', 'batch_image_token'])
        product.invalidate_recordset(['image_1920'])
        self.env['ir.attachment'].invalidate_model(['checksum', 'datas'])
        replacement_attachment = self._piece_image_attachments(
            line
        ).ensure_one()
        self.assertEqual(line.batch_image_1920, TEST_IMAGE_2)
        self.assertTrue(line.batch_image_token)
        self.assertNotEqual(line.batch_image_token, first_token)
        self.assertNotEqual(replacement_attachment.checksum, first_checksum)
        self.assertEqual(product.image_1920, TEST_IMAGE)

    def test_direct_piece_image_upload_rejects_stale_image_snapshot(self):
        fixture = self._create_fixture()
        line = fixture['lines'][0]
        first_row = self._open_setup(
            fixture['production']
        ).line_ids.filtered(
            lambda row: row.production_line_id == line
        ).ensure_one()
        stale_row = self._open_setup(
            fixture['production']
        ).line_ids.filtered(
            lambda row: row.production_line_id == line
        ).ensure_one()

        self.assertTrue(first_row.upload_piece_image(TEST_IMAGE))
        line.invalidate_recordset(['batch_image_1920', 'batch_image_token'])
        saved_token = line.batch_image_token

        with self.assertRaises(UserError):
            stale_row.upload_piece_image(TEST_IMAGE_2)

        line.invalidate_recordset(['batch_image_1920', 'batch_image_token'])
        self.assertEqual(line.batch_image_1920, TEST_IMAGE)
        self.assertEqual(line.batch_image_token, saved_token)
        self.assertEqual(len(self._piece_image_attachments(line)), 1)

    def test_direct_piece_image_upload_rejects_invalid_payloads(self):
        fixture = self._create_fixture()
        line = fixture['lines'][0]
        setup_line = self._open_setup(
            fixture['production']
        ).line_ids.filtered(
            lambda row: row.production_line_id == line
        ).ensure_one()
        fake_webp = base64.b64encode(
            b'RIFF\x18\x00\x00\x00WEBPVP8 ' + (b'\x00' * 24)
        )
        invalid_payloads = (
            b'not-valid-base64!!!',
            base64.b64encode(b'%PDF-1.4 not an image'),
            fake_webp,
            'data:image/jpeg;base64,%s' % TEST_IMAGE.decode(),
        )

        for payload in invalid_payloads:
            with self.assertRaises(ValidationError):
                setup_line.upload_piece_image(payload)

        with patch.object(
            mrp_tailoring_material_setup,
            'PIECE_IMAGE_MAX_UPLOAD_BYTES',
            8,
        ):
            with self.assertRaisesRegex(ValidationError, 'حجم'):
                setup_line.upload_piece_image(base64.b64encode(b'x' * 9))

        line.invalidate_recordset(['batch_image_1920', 'batch_image_token'])
        self.assertFalse(line.batch_image_1920)
        self.assertFalse(line.batch_image_token)
        self.assertFalse(self._piece_image_attachments(line))

    def test_direct_piece_image_upload_enforces_user_and_row_scope(self):
        fixture = self._create_fixture()
        line = fixture['lines'][0]
        regular_setup = self.env[
            'furniture.mrp.tailoring.setup.wizard'
        ].with_user(self.regular_user).create({
            'production_id': fixture['production'].id,
            'line_ids': [(0, 0, {
                'sequence': line.sequence,
                'production_line_id': line.id,
            })],
        })
        regular_setup_line = regular_setup.line_ids.ensure_one()

        with self.assertRaises(AccessError):
            regular_setup_line.upload_piece_image(TEST_IMAGE_2)

        setup = self._open_setup(fixture['production'])
        setup_line = setup.line_ids.filtered(
            lambda row: row.production_line_id == line
        ).ensure_one()

        foreign = self._create_fixture()
        foreign_line = foreign['lines'][0]
        setup_line.write({'production_line_id': foreign_line.id})
        with self.assertRaises(AccessError):
            setup_line.upload_piece_image(TEST_IMAGE_2)

        line.invalidate_recordset(['batch_image_1920', 'batch_image_token'])
        foreign_line.invalidate_recordset([
            'batch_image_1920',
            'batch_image_token',
        ])
        self.assertFalse(line.batch_image_1920)
        self.assertFalse(foreign_line.batch_image_1920)

        fresh_setup = self._open_setup(fixture['production'])
        expired_row = fresh_setup.line_ids.filtered(
            lambda row: row.production_line_id == line
        ).ensure_one()
        expired_row_id = expired_row.id
        expired_row.unlink()
        with self.assertRaises(UserError):
            self.env[
                'furniture.mrp.tailoring.setup.wizard.line'
            ].browse(expired_row_id).upload_piece_image(TEST_IMAGE_2)

    def test_direct_piece_image_upload_rejects_terminal_orders(self):
        for terminal_state in ('done', 'cancelled'):
            fixture = self._create_fixture()
            line = fixture['lines'][0]
            setup_line = self._open_setup(
                fixture['production']
            ).line_ids.filtered(
                lambda row: row.production_line_id == line
            ).ensure_one()
            fixture['production'].with_context(
                furniture_skip_material_refresh=True,
                furniture_skip_stage_plan_sync=True,
            ).write({'state': terminal_state})

            with self.assertRaises(UserError):
                setup_line.upload_piece_image(TEST_IMAGE_2)
            line.invalidate_recordset([
                'batch_image_1920',
                'batch_image_token',
            ])
            self.assertFalse(line.batch_image_1920)
            self.assertFalse(line.batch_image_token)

    def test_order_piece_images_are_cleared_on_done_and_cancel(self):
        for terminal_state in ('done', 'cancelled'):
            fixture = self._create_fixture()
            production = fixture['production']
            line = fixture['lines'][0]
            product = line.product_id
            setup = self._open_setup(production)
            setup_line = setup.line_ids.filtered(
                lambda row: row.production_line_id == line
            ).ensure_one()
            self.assertTrue(setup_line.upload_piece_image(TEST_IMAGE_2))
            self.assertTrue(setup.add_order_image(TEST_IMAGE, 'main.png'))
            self.assertTrue(setup.add_order_image(TEST_IMAGE_2, 'extra.png'))
            line.invalidate_recordset([
                'batch_image_1920',
                'batch_image_token',
            ])
            production.invalidate_recordset([
                'tailoring_set_image_1920',
                'tailoring_set_image_token',
            ])
            self.assertTrue(line.batch_image_1920)
            self.assertTrue(line.batch_image_token)
            self.assertTrue(production.tailoring_set_image_1920)
            self.assertTrue(production.tailoring_set_image_token)
            self.assertEqual(
                len(production.sudo().tailoring_set_image_attachment_ids),
                1,
            )
            self.assertEqual(len(self._piece_image_attachments(line)), 1)
            self.assertEqual(len(self._order_image_attachments(production)), 1)

            production.with_context(
                furniture_skip_material_refresh=True,
                furniture_skip_stage_plan_sync=True,
            ).write({'state': terminal_state})

            line.invalidate_recordset([
                'batch_image_1920',
                'batch_image_token',
            ])
            production.invalidate_recordset([
                'tailoring_set_image_1920',
                'tailoring_set_image_token',
            ])
            product.invalidate_recordset(['image_1920'])
            self.assertFalse(line.batch_image_1920)
            self.assertFalse(line.batch_image_token)
            self.assertFalse(production.tailoring_set_image_1920)
            self.assertFalse(production.tailoring_set_image_token)
            self.assertFalse(
                production.sudo().tailoring_set_image_attachment_ids
            )
            self.assertFalse(self._piece_image_attachments(line))
            self.assertFalse(self._order_image_attachments(production))
            self.assertEqual(product.image_1920, TEST_IMAGE)

    def test_ordinary_copy_drops_image_but_operational_split_preserves_it(self):
        fixture = self._create_fixture(line_quantities=(4.0, 1.0, 1.0))
        line = fixture['lines'][0]
        line.write({
            'batch_image_1920': TEST_IMAGE_2,
            'batch_image_token': 'piece-operational-split',
        })

        ordinary_copy = line.with_context(
            furniture_skip_line_consolidation=True,
            furniture_skip_material_refresh=True,
            furniture_skip_stage_plan_sync=True,
            furniture_skip_tailoring_allocation_copy=True,
        ).copy({
            'sequence': line.sequence + 1,
            'product_qty': 1.0,
        })
        self.assertFalse(ordinary_copy.batch_image_1920)
        self.assertFalse(ordinary_copy.batch_image_token)

        remaining = line._split_for_partial_quantity(1.5)
        self.assertTrue(remaining)
        line.invalidate_recordset(['batch_image_1920', 'batch_image_token'])
        remaining.invalidate_recordset([
            'batch_image_1920',
            'batch_image_token',
        ])
        self.assertEqual(line.batch_image_1920, TEST_IMAGE_2)
        self.assertEqual(remaining.batch_image_1920, TEST_IMAGE_2)
        self.assertEqual(
            remaining.batch_image_token,
            line.batch_image_token,
        )
